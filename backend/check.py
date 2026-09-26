"""Reproducible checks: the default replaces only the external model endpoint."""
import os
import sys
import time
import io
from pathlib import Path
from docx import Document
from .app import Service
from .check_harness import main as contract_check


def real_check():
    if not os.getenv('DEEPSEEK_API_KEY'):
        raise SystemExit('DEEPSEEK_API_KEY is missing; no model call attempted')
    pdf_dir = Path(os.getenv('METHODATLAS_PDF_DIR', '论文'))
    data_dir = Path(os.getenv('METHODATLAS_CHECK_DATA', '.methodatlas-data/real-check'))
    service = Service(pdf_dir,data_dir)
    try:
        project = service.store.projects()[-1]
        papers = service.store.papers(project['id'])
        chosen = [p['id'] for p in papers if '04_Diffusion_Policy' in p['title'] or '10_DP3_' in p['title']]
        if len(chosen) != 2:
            raise SystemExit('The real check needs Diffusion Policy and DP3 PDFs')
        conversation = os.getenv('METHODATLAS_CHECK_CONVERSATION') or service.store.create_conversation(project['id'])
        print(f'Validation conversation: {conversation}', flush=True)
        files, saved = {}, {}
        for index, prompt in enumerate([
            '仅基于勾选的两篇论文，简述输入表征和策略生成的区别，并给原文依据。',
            '刚才 DP3 的策略骨干沿用了什么？两句话回答。',
            '在 DP3 中找出并高亮稀疏点云编码和沿用 CNN 扩散骨干的多处原文。',
            '把这些部分写成简短中文综述 DOCX，保留实际依据。',
            '给刚才的综述增加一个对比表，交付新版 DOCX。',
            '另产一份 HTML，用流程图可视化刚才的方法区别，原文引用能点击。',
            '把刚才的 HTML 图表颜色改成蓝绿色，保留原文引用。',
            '仅基于勾选的验收配色笔记，告诉我感知和动作分别用什么颜色，并引用笔记；不要生成文件。',
            '刚才笔记给动作约定了什么颜色？请回读笔记并引用，一句话回答，不生成文件。',
        ]):
            if index == 7:
                title = f'验收配色笔记 {conversation}'
                note = next((p['id'] for p in service.store.papers(project['id']) if p['title'] == title), None)
                chosen = [note or service.store.paste(project['id'], title, '个人展示约定：感知使用青色，动作使用紫色。')]
            if index == 8:
                before = service.project_state(project['id'])
                service.close()
                service = Service(pdf_dir, data_dir)
                assert service.project_state(project['id']) == before, 'Restart changed saved project state'
            body = {'text':prompt,'selected_paper_ids':chosen,'client_message_id':f'{conversation}-{index}'}
            submitted = service.message(project['id'],conversation,body)
            while True:
                task = service.task_view(project['id'],submitted['task_id'])
                if task['status'] in ('succeeded','failed','interrupted'):
                    break
                time.sleep(.5)
            assert task['status'] == 'succeeded', task.get('error')
            if index < 3 or index >= 7:
                assert not task['files']
            else:
                assert len(task['files']) == 1
                file = task['files'][0]
                kind = 'docx' if index < 5 else 'html'
                assert file['kind'] == kind
                version = service.file_version(project['id'], file['artifact_id'], file['version_id'], conversation)
                raw = service.file_path(project['id'], version).read_bytes()
                if index in (4, 6):
                    assert file['artifact_id'] == files[kind]['artifact_id'], 'Revision created a separate artifact'
                    assert file['version_no'] == 2 and version['payload']['base_version_id'] == files[kind]['version_id']
                else:
                    assert file['version_no'] == 1
                    files[kind] = file
                if kind == 'docx':
                    document = Document(io.BytesIO(raw))
                    assert document.paragraphs
                    if index == 4:
                        assert document.tables, 'Requested comparison table is absent'
                else:
                    assert b'<html' in raw and b'Content-Security-Policy' in raw
                saved[file['version_id']] = (file['artifact_id'], raw)
            if index != 1:
                assert task['citations'], 'Expected source evidence'
            if index == 2:
                assert len({c['id'] for c in task['citations'] if c['rect']}) >= 2
            if index >= 7:
                assert {c['paper_id'] for c in task['citations']} == set(chosen)
            if index == 8:
                assert any(t['name'] in ('read_material', 'read_evidence', 'locate') and t['status'] == 'succeeded'
                           for t in task['tools']), 'Restart follow-up did not actually reread evidence'
            assert service.message(project['id'],conversation,body)['duplicate']
            for version_id, (artifact_id, raw) in saved.items():
                version = service.file_version(project['id'], artifact_id, version_id, conversation)
                assert service.file_path(project['id'], version).read_bytes() == raw, 'An old file changed'
            print(f'PASS real step {index + 1}: {len(task["citations"])} citations, {len(task["files"])} files',flush=True)
        print(f'Preserved validation data: {data_dir.resolve()}')
    finally:
        service.close()


if __name__ == '__main__':
    real_check() if '--real' in sys.argv[1:] else contract_check()
