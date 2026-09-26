"""External HTTP regression check for #12 (no provider credentials required)."""
import json
import sqlite3
import io
import tempfile
import threading
import time
import urllib.error
import urllib.request
from collections import deque
from contextlib import contextmanager, nullcontext
from unittest.mock import patch
from http.server import BaseHTTPRequestHandler
from http.server import ThreadingHTTPServer
from pathlib import Path
from types import SimpleNamespace

from .app import Handler, Service
from .files import render_file
from .paper_research import bound_card, verify_claims
from .research_context import References, model_view
from .state import json_text
from concurrent.futures import ThreadPoolExecutor
from docx import Document
import pymupdf


def check_research_budgets():
    citation = {'id':'cite_' + '1' * 32, 'paper_id':'paper', 'paper_version_id':'version',
                'page':1, 'block':1, 'quote':'x' * 20000}
    calls = []

    def complete(task, system, context, role, **kwargs):
        assert role == 'verify' and len(json_text(context)) <= 60000
        calls.append([item['index'] for item in context['claims']])
        return {'verdicts':[{'index':item['index'], 'supported':item['index'] % 2 == 0,
                             'reason':str(item['index'])} for item in reversed(context['claims'])]}

    db = sqlite3.connect(':memory:')
    db.row_factory = sqlite3.Row
    db.execute('CREATE TABLE claim_checks(project_id TEXT,cache_key TEXT,result TEXT,PRIMARY KEY(project_id,cache_key))')
    store = SimpleNamespace(citation=lambda *args: citation,
                            one=lambda sql,args: db.execute(sql,args).fetchone(),
                            transaction=lambda: nullcontext(db),
                            assert_active=lambda *args: None)
    tools = SimpleNamespace(task={'id':'budget','revision':0}, store=store,project='project',conversation='conversation', references=References(),
                            read_evidence=lambda ids: {'evidence':[citation for _ in ids]},
                            research=SimpleNamespace(complete=complete))
    claims = [{'text':f'Claim {index}', 'citation_ids':[citation['id']]} for index in range(4)]
    verdicts = verify_claims(tools, claims)['verdicts']
    assert calls == [[0, 1, 2, 3]]
    assert [v['index'] for v in verdicts] == list(range(4))
    assert [v['supported'] for v in verdicts] == [True, False, True, False]
    assert [v['reason'] for v in verdicts] == ['0', '1', '2', '3']

    calls.clear()
    citation['quote'] = 'x' * 60001
    verdicts = verify_claims(tools, claims[:1])['verdicts']
    assert not calls and len(verdicts) == 1
    assert verdicts[0]['index'] == 0 and verdicts[0]['supported'] is False
    assert '尚未核验' in verdicts[0]['reason']

    original = [{'text':f'{index}: ' + 'x' * 1500, 'citation_ids':[citation['id']]} for index in range(32)]
    card = bound_card({'claims':list(original), 'gaps':['Known evidence gap.']})
    assert len(json_text(card)) <= 18000 and 0 < len(card['claims']) < len(original)
    assert card['claims'] == original[:len(card['claims'])]
    assert card['gaps'] == ['Known evidence gap.']
    assert card['budget_omissions'] == len(original) - len(card['claims'])
    assert '省略' in card['budget_note'] and '不能视为完整覆盖' in card['budget_note']
    prior_omissions = card['budget_omissions']
    card['claims'].extend(original[-3:])
    before_second_trim = len(card['claims'])
    bound_card(card)
    assert len(card['claims']) < before_second_trim and len(json_text(card)) <= 18000
    assert card['budget_omissions'] == prior_omissions + before_second_trim - len(card['claims'])
    trimmed = json_text(card)
    assert json_text(bound_card(card)) == trimmed
    print('PASS: in-memory verification budgets, preserved indices, oversized evidence and explicit card omissions')


def main():
    check_research_budgets()
    references = References()
    ids = ['cite_' + f'{i:032x}' for i in range(200)]
    with ThreadPoolExecutor(max_workers=12) as pool:
        aliases = list(pool.map(references.alias, ids * 5))
    assert len(set(aliases)) == 200
    assert [references.resolve(alias) for alias in aliases] == ids * 5
    refs = [{'id':'cite_' + str(i) * 32, 'title':title, 'paper_version_id':title, 'page':1, 'rect':None, 'quote':title + ' evidence'} for i,title in [(1,'A'),(2,'B')]]
    source = f'A claim [cite:{refs[0]["id"]}]. B claim [cite:{refs[1]["id"]}].'
    raw, preview = render_file('docx','Stable references',source,list(reversed(refs)))
    assert preview == 'A claim [2]. B claim [1].'
    assert Document(io.BytesIO(raw)).paragraphs[1].text == preview
    controlled = {**refs[0], 'quote':'原文\x00内容\x0b\n第二行'}
    raw, preview = render_file('docx', '标题\x00', '正文\x0c', [controlled])
    exported = '\n'.join(p.text for p in Document(io.BytesIO(raw)).paragraphs)
    assert '原文内容' not in exported and '[1] A' in exported and preview == '正文'
    assert controlled['quote'] == '原文\x00内容\x0b\n第二行'
    repeated = [refs[0], {**refs[0], 'id': 'cite_' + '3' * 32}]
    for kind, content in [('docx', 'Claim'), ('html', '<p>Claim</p>')]:
        raw, _ = render_file(kind, 'Sources', content, repeated)
        text = raw.decode() if kind == 'html' else '\n'.join(p.text for p in Document(io.BytesIO(raw)).paragraphs)
        if kind == 'docx':
            assert text.count('[1, 2] A') == 1 and 'A evidence' not in text
        else:
            assert all(f'id="ref-{c["id"]}"' in text for c in repeated)
            assert 'A evidence' in text

    for invalid, supplied in [('Claim [1].', refs), ('Claim [1].', []), ('Claim [cite:unknown].', refs), ('Claim [cite:unfinished', refs)]:
        try:
            render_file('docx','Invalid references',invalid,supplied)
        except ValueError:
            pass
        else:
            raise AssertionError('Invalid DOCX citation accepted')
    # Windows cannot always remove the SDK's deeply nested profile links.
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
        root = Path(directory)
        (root / 'papers').mkdir()
        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((60,80), 'Original evidence: blue method uses point clouds.')
        page.insert_text((60,150), 'Second evidence: evaluation is indoors only.')
        pdf.save(root / 'papers' / 'check.pdf')
        pdf.close()
        plans = deque()

        class Provider(BaseHTTPRequestHandler):
            def log_message(self, *_):
                pass

            def do_POST(self):
                body = json.loads(self.rfile.read(int(self.headers['Content-Length'])))
                assert all(t['function']['name'].startswith('mcp__methodatlas__') for t in body.get('tools', [])), 'Unexpected shell or other model tools'
                # Existing cases exercise Agent/tool contracts; route them explicitly.
                planning = any('默认mode="rag"' in str(m.get('content', '')) for m in body['messages'])
                pipeline_case = isinstance(plans[0], dict) and plans[0].get('_rag', False)
                if planning and not pipeline_case:
                    step = {'role': 'assistant', 'content': json.dumps({'intent':'research','mode':'explore', 'reason':'测试探索工具契约', 'only_selected':False})}
                else:
                    step = plans.popleft()
                    if isinstance(step, dict):
                        step.pop('_rag', None)
                message = step(body) if callable(step) else step
                assert (self.headers.get('x-deepseek-harness-compact') == '1') == message.pop('_compact', False)
                usage = message.pop('_usage', None)
                if '_http_error' in message:
                    raw = json.dumps({'error':{'message':message['_http_error']['message'],'type':'insufficient_quota','code':'QUOTA'}}).encode()
                    self.send_response(message['_http_error']['status'])
                    self.send_header('Content-Type','application/json')
                    self.send_header('Content-Length',str(len(raw)))
                    self.end_headers()
                    self.wfile.write(raw)
                    return
                finish = 'tool_calls' if message.get('tool_calls') else 'stop'
                chunks = [{'id':'local-check','object':'chat.completion.chunk','created':0,'model':body['model'], 'choices':[{'index':0,'delta':message,'finish_reason':None}]},
                          {'id':'local-check','object':'chat.completion.chunk','created':0,'model':body['model'], 'choices':[{'index':0,'delta':{},'finish_reason':finish}]}]
                if usage:
                    chunks[-1]['usage'] = usage
                raw = ''.join('data: ' + json.dumps(chunk) + '\n\n' for chunk in chunks).encode() + b'data: [DONE]\n\n'
                self.send_response(200)
                self.send_header('Content-Type','text/event-stream')
                self.send_header('Content-Length',str(len(raw)))
                self.end_headers()
                self.wfile.write(raw)

        provider = ThreadingHTTPServer(('127.0.0.1',0),Provider)
        threading.Thread(target=provider.serve_forever,daemon=True).start()
        service = Service(root / 'papers', root / 'data')
        service.research.key = 'local-test-only'
        service.research.base_url = f'http://127.0.0.1:{provider.server_port}'
        handler = type('CheckHandler', (Handler,), {'service': service})
        server = ThreadingHTTPServer(('127.0.0.1', 0), handler)
        threading.Thread(target=server.serve_forever, daemon=True).start()
        base = f'http://127.0.0.1:{server.server_port}'

        def request(path, body=None):
            req = urllib.request.Request(base + path,
                data=json.dumps(body).encode() if body is not None else None,
                headers={'Content-Type': 'application/json', 'Origin': base})
            with urllib.request.urlopen(req, timeout=10) as response:
                return json.load(response)

        def binary(path):
            with urllib.request.urlopen(base + path) as response:
                return response.read()

        def rejected(path, body=None, code=400):
            try:
                request(path,body)
            except urllib.error.HTTPError as error:
                assert error.code == code
            else:
                raise AssertionError('invalid ownership/input was accepted')

        def tool(name, arguments):
            return {'role':'assistant','tool_calls':[{'index':0,'id':'call_' + str(time.time_ns()),'type':'function','function':{'name':'mcp__methodatlas__' + name,'arguments':json.dumps(arguments)}}]}

        def final(text='已完成。'):
            return {'role':'assistant','content':text}

        def submit(text, key, expected='succeeded', report_sources=None):
            body = {'text':text,'client_message_id':key,'selected_paper_ids':[pdf_id]}
            if report_sources is not None:
                body['report_sources'] = report_sources
            response = request(message_path,body)
            task_path = f'/api/projects/{project["id"]}/tasks/{response["task_id"]}'
            deadline = time.monotonic() + 90
            while time.monotonic() < deadline:
                task = request(task_path)
                if task['status'] in ('succeeded','failed','interrupted'):
                    assert task['status'] == expected, task.get('error')
                    return body, response, task
                time.sleep(.1)
            raise AssertionError('Harness did not complete the local provider check')

        try:
            project = request('/api/state')['projects'][0]
            pdf_id = project['papers'][0]['id']
            conversation = project['conversations'][0]['id']
            message_path = f'/api/projects/{project["id"]}/conversations/{conversation}/messages'
            source = request(f'/api/projects/{project["id"]}/materials',
                             {'title': '研究笔记', 'text': '蓝色方法只在室内验证。'})
            paper = request(f'/api/projects/{project["id"]}/papers/{source["paper_id"]}')
            assert paper['pages'][0]['text'] == '蓝色方法只在室内验证。'
            print('PASS: pasted material persists with a source version')
            raced_pdf = root / 'replaced-during-parse.pdf'
            raced_pdf.write_bytes((root / 'papers' / 'check.pdf').read_bytes())
            replacement = pymupdf.open()
            replacement.new_page().insert_text((60,80), 'Replacement bytes must not inherit original evidence.')
            replacement_bytes = replacement.tobytes()
            replacement.close()
            open_pdf = pymupdf.open
            @contextmanager
            def replace_after_parse(*args, **kwargs):
                with open_pdf(*args, **kwargs) as parsed:
                    yield parsed
                raced_pdf.write_bytes(replacement_bytes)
            with patch('backend.pdf.pymupdf.open', replace_after_parse):
                try:
                    service.store.import_pdf(project['id'], raced_pdf)
                except ValueError as error:
                    assert '改变' in str(error)
                else:
                    raise AssertionError('evidence was bound to replacement PDF bytes')
            plans.extend([{**final(json.dumps({'intent':'chat','mode':'rag','only_selected':True,'reads':[], 'files':[]})), '_rag':True},
                          final(json.dumps({'answer':'你好。', 'files':[]}))])
            ordinary, response, task = submit('你好', 'ordinary')
            assert task['plan'] == []
            assert not task['files'] and not request(f'/api/projects/{project["id"]}')['artifacts']
            assert request(message_path,ordinary)['task_id'] == response['task_id']
            rejected(message_path,{**ordinary,'text':'changed'})
            plans.append({**final('你好！有什么可以帮你？'), '_rag':True})
            _, _, direct = submit('你好', 'direct-greeting')
            assert direct['kind'] == 'chat' and not direct['tools'] and not direct['files'] and not direct['plan']
            assert not plans, 'a simple greeting must complete in one model call'
            plans.extend([tool('set_scope',{'only_selected':True}), tool('read_material',{'paper_id':source['paper_id']}), final('范围外材料未读取。')])
            _, _, limited = submit('仅限勾选材料', 'limited')
            assert not limited['citations'] and not limited['files']
            assert any(t['name'] == 'read_material' and t['status'] == 'failed' for t in limited['tools'])
            plans.extend([tool('set_scope',{'only_selected':False}), tool('read_material',{'paper_id':source['paper_id']}), final('允许使用项目内其他材料。')])
            _, _, expanded = submit('可以参考我的项目笔记', 'expanded')
            assert expanded['citations'][0]['paper_id'] == source['paper_id'] and not expanded['files']
            # #14: the public Harness boundary must expose continuation and real read ranges.
            def check_first_page(body):
                result = json.loads(body['messages'][-1]['content'])
                if 'evidence' not in result:
                    result = json.loads(result['content'][0]['text'])
                assert result['next_offset'] == 1 and result['matched_blocks'] == 2
                assert 'coverage' not in result and 'rect' not in result['evidence'][0]
                assert result['evidence'][0]['id'].startswith('E')
                return tool('read_material', {'paper_id':pdf_id, 'limit':1, 'offset':result['next_offset']})
            plans.extend([tool('set_scope',{'only_selected':True}),
                          tool('read_material',{'paper_id':pdf_id,'limit':1}), check_first_page, final('分段读取了两处正文。')])
            _, _, paged = submit('分段读取勾选论文', 'paged')
            assert len(paged['citations']) == 2
            assert paged['coverage'][paged['citations'][0]['paper_version_id']]['read_blocks'] == 2
            initial_plan = [{'id':'read','title':'读取相关材料','status':'in_progress'}, {'id':'answer','title':'整理结论','status':'pending'}]
            next_plan = [{'id':'read','title':'读取相关材料','status':'completed','summary':'已读取项目笔记中的展示约定。'}, {'id':'answer','title':'整理结论','status':'in_progress'}]
            complete_plan = [next_plan[0], {'id':'answer','title':'整理结论','status':'completed','summary':'已根据笔记整理回答，未生成文件。'}]
            plans.extend([tool('update_plan',{'steps':initial_plan}), tool('set_scope',{'only_selected':False}), tool('read_material',{'paper_id':source['paper_id']}),
                          tool('update_plan',{'steps':next_plan}), tool('update_plan',{'steps':complete_plan}), final('已经按笔记整理完成。')])
            _, planned_response, planned = submit('读取笔记并核对展示约定后整理回答', 'planned')
            assert planned['plan'] == complete_plan and not planned['files']
            grouped_events = [e for e in planned['events'] if e['status'].startswith('tool_')]
            assert grouped_events and all(e['step_id'] == 'read' for e in grouped_events)
            assert any('阅读' in e['message'] and '研究笔记' in e['message'] for e in grouped_events)
            plans.extend([tool('update_plan',{'steps':[{'id':'x','title':'A','status':'in_progress'},{'id':'y','title':'B','status':'in_progress'}]}),
                          tool('update_plan',{'steps':[{'id':'x','title':'A','status':'completed','summary':'未执行'}]}), final('无效计划未保存。')])
            _, _, invalid_plan = submit('计划边界检查', 'invalid-plan')
            assert not invalid_plan['plan'] and sum(t['status'] == 'failed' for t in invalid_plan['tools']) == 2
            plans.extend([tool('update_plan',{'steps':initial_plan}),
                          tool('update_plan',{'steps':[{**initial_plan[0], 'title':'改写已执行的工作'},initial_plan[1]]}),
                          tool('update_plan',{'steps':[{**initial_plan[0], 'status':'pending'},initial_plan[1]]}),
                          final('已开始步骤的名称和状态均保留。')])
            _, _, preserved_plan = submit('保留步骤记录', 'preserved-plan')
            assert preserved_plan['plan'] == initial_plan and sum(t['status'] == 'failed' for t in preserved_plan['tools']) == 2
            plans.extend([tool('set_scope',{'only_selected':True}), tool('locate',{'paper_id':pdf_id,'query':'evidence'})])
            def write_from_evidence(body):
                evidence = json.loads(body['messages'][-1]['content'])
                # Harness MCP projection may wrap the tool's text; locate the JSON text.
                if 'evidence' not in evidence and 'content' in evidence:
                    evidence = json.loads(evidence['content'][0]['text'])
                citation_ids[:] = [c['id'] for c in evidence['evidence']]
                arguments = {'title':'Check document','kind':'docx','content':f'# Finding\nBlue method uses point clouds. [cite:{citation_ids[0]}]','citation_ids':citation_ids}
                first = tool('write_file', arguments)
                duplicate = tool('write_file', {**arguments, 'content':arguments['content'].replace('[cite:cite_', '[cite:'), 'citation_ids':[cid.removeprefix('cite_') for cid in citation_ids]})['tool_calls'][0]
                duplicate['index'] = 1
                first['tool_calls'].append(duplicate)
                return first
            citation_ids = []
            plans.extend([write_from_evidence,final()])
            generated, response, task = submit('生成 DOCX', 'generated')
            assert len(task['files']) == 1 and len(task['citations']) == 2
            citation_ids[:] = [c['id'] for c in task['citations']]
            file = task['files'][0]
            prefix = f'/api/projects/{project["id"]}/artifacts/{file["artifact_id"]}'
            old_url = prefix + f'/versions/{file["version_id"]}/download'
            original = binary(old_url)
            assert Document(io.BytesIO(original)).paragraphs[1].text == 'Finding'
            assert request(message_path,generated)['duplicate']
            c = task['citations'][0]
            assert c['page'] == 1 and c['rect'][0] == 60
            short_id = c['id'].removeprefix('cite_')
            plans.append(final(f'已核对原文。[cite:{short_id}] 再次引用。[cite:{c["id"]}]'))
            _, _, cited = submit('保留可点击的原文链接', 'short-citation')
            saved_answer = request(f'/api/projects/{project["id"]}')['conversations'][0]['messages'][-1]['text']
            assert saved_answer.count(f'[cite:{c["id"]}]') == 2 and [item['id'] for item in cited['citations']] == [c['id']]
            assert request(f'/api/projects/{project["id"]}/citations/{short_id}')['id'] == c['id']
            other = request('/api/projects',{'name':'Other'})['project_id']
            rejected(f'/api/projects/{other}/citations/{c["id"]}')
            rejected(f'/api/projects/{other}/citations/{short_id}')
            def repaired_answer(body):
                feedback = json.loads(body['messages'][-1]['content'])
                assert [e['id'] for e in feedback['lines'][0]['errors']] == ['id', 'missing']
                assert 'UNCHANGED' not in json.dumps(feedback) and not body.get('tools')
                assert all(e['paper_id'] == pdf_id for e in feedback['candidates'])
                return final(json.dumps({'replacements':[{'line':1,'text':f'已纠正。[cite:{short_id}]'}]}))
            plans.extend([tool('set_scope', {'only_selected':True}), final(f'UNCHANGED\n有效。[cite:{short_id}] 错误。[cite:id][cite:missing]'), repaired_answer])
            _, _, repaired = submit('纠正回答中的错误引用', 'repair-citation')
            assert [item['id'] for item in repaired['citations']] == [c['id']]
            assert sum(e['status'] == 'citation_repair' for e in repaired['events']) == 1
            assert len(service.store.all("SELECT id FROM messages WHERE task_id=? AND role='assistant'", (repaired['id'],))) == 1
            assert not repaired['files']
            assert service.store.one("SELECT text FROM messages WHERE task_id=? AND role='assistant'", (repaired['id'],))['text'].startswith('UNCHANGED\n')
            def invalid_repair(text):
                return final(json.dumps({'replacements':[{'line':0,'text':text}]}))
            plans.extend([final('[cite:' + '0' * 32 + ']'), invalid_repair('[cite:' + '0' * 32 + ']')])
            _, _, invalid_citation = submit('不存在的引用不能保存', 'unknown-citation', expected='failed')
            assert not invalid_citation['citations']
            original_message_path = message_path
            second_conversation = request(f'/api/projects/{project["id"]}/conversations',{})['conversation_id']
            message_path = f'/api/projects/{project["id"]}/conversations/{second_conversation}/messages'
            plans.extend([final(f'[cite:{short_id}]'), invalid_repair(f'[cite:{short_id}]')])
            _, _, foreign_citation = submit('不能引用其他对话证据', 'foreign-citation', expected='failed')
            assert not foreign_citation['citations']
            message_path = original_message_path
            outside = f'[cite:{expanded["citations"][0]["id"].removeprefix("cite_")}]'
            plans.extend([tool('set_scope',{'only_selected':True}), final(outside), invalid_repair(outside)])
            _, _, outside_scope = submit('不能引用范围外材料', 'scope-citation', expected='failed')
            assert outside_scope['error_info']['code'] == 'evidence' and not outside_scope['citations']
            rejected(f'/api/projects/{other}/papers/{pdf_id}?version_id={c["paper_version_id"]}')
            rejected(f'/api/projects/{other}/artifacts/{file["artifact_id"]}/versions/{file["version_id"]}/download')
            plans.extend([tool('set_scope',{'only_selected':True}), tool('read_file',{'artifact_id':file['artifact_id'],'version_id':file['version_id']}),
                          tool('write_file',{'title':'Check document','kind':'docx','content':f'Shorter. [cite:{citation_ids[0]}]\n\n| Method | Input |\n|---|---|\n| Blue | Point cloud |','citation_ids':citation_ids,'artifact_id':file['artifact_id'],'base_version_id':file['version_id']}),final()])
            _, _, task = submit('增加对比表', 'revision')
            assert task['files'][0]['artifact_id'] == file['artifact_id']
            assert task['files'][0]['downloadable'] and len(task['citations']) == len(citation_ids)
            read_result = json.loads(service.store.one("SELECT result FROM tool_calls WHERE task_id=? AND name='read_file'", (task['id'],))['result'])
            assert f'[cite:{citation_ids[0]}]' in read_result['body']
            versions = request(prefix)['versions']
            assert len(versions) == 2 and binary(old_url) == original
            latest = binary(prefix + f'/versions/{versions[-1]["id"]}/download')
            assert len(Document(io.BytesIO(latest)).tables) == 1
            plans.extend([tool('set_scope',{'only_selected':True}), tool('read_file',{'artifact_id':file['artifact_id'],'version_id':file['version_id']}),
                          tool('write_file',{'title':'Stale write','kind':'docx','content':'Must not replace a newer version.','citation_ids':citation_ids,'artifact_id':file['artifact_id'],'base_version_id':file['version_id']}), final('基线已变化，没有写入。')])
            _, _, failed_write = submit('错误基线写入检查', 'stale')
            assert not failed_write['files'] and len(request(prefix)['versions']) == 2 and binary(old_url) == original
            plans.extend([tool('set_scope',{'only_selected':True}), tool('write_file',{'title':'Invalid HTML citation','kind':'html','content':'<button data-citation="' + '0' * 32 + '">Unknown</button>', 'citation_ids':citation_ids}), final('无效 HTML 引用未写入。')])
            _, _, invalid_html = submit('拒绝未知 HTML 引用', 'invalid-html-citation')
            assert not invalid_html['files'] and any(t['name'] == 'write_file' and t['status'] == 'failed' for t in invalid_html['tools'])
            html_arguments = {'title':'HTML check','kind':'html','content':f'<h1>Blue</h1><svg viewBox="0 0 100 100"><rect width="50" height="50" fill="blue"/></svg><button data-citation="{short_id}">原文</button><script>fetch("https://evil.invalid")</script><img src="https://evil.invalid"><a href="https://evil.invalid">unsafe</a>', 'citation_ids':[short_id,*citation_ids[1:]]}
            css = '.flow > [class="node"] { font-family: "Segoe UI"; color: teal; }'
            html_arguments['content'] += f'<style>{css}</style><div class="flow"><span class="node">Styled diagram</span></div>'
            plans.extend([tool('set_scope',{'only_selected':True}), tool('write_file',html_arguments),
                          tool('write_file',{**html_arguments, 'citation_ids':citation_ids, 'content':html_arguments['content'].replace(f'data-citation="{short_id}"', f'data-citation="{c["id"]}"')}), final()])
            _, _, html_task = submit('生成 HTML', 'html')
            assert len(html_task['files']) == 1
            html_file = html_task['files'][0]
            html_url = f'/api/projects/{project["id"]}/artifacts/{html_file["artifact_id"]}/versions/{html_file["version_id"]}/download'
            html_bytes = binary(html_url)
            assert css.encode() in html_bytes, 'Downloaded HTML must preserve CSS quotes and child selectors'
            assert f'data-citation="{c["id"]}"'.encode() in html_bytes and f'data-citation="{short_id}"'.encode() not in html_bytes
            assert b'evil.invalid' not in html_bytes and b'Content-Security-Policy' in html_bytes and b'postMessage' in html_bytes
            assert b'<section class="methodatlas-sources">' in html_bytes
            html_version = service.store.artifact(project['id'], html_file['artifact_id'])['versions'][-1]
            view = model_view('read_file', html_version, References())
            assert 'Original evidence:' not in view['body'] and '原文依据' not in view['body']
            assert all('quote' not in item for item in view['citations'])
            assert b'Original evidence:' in html_bytes
            # Fail the material reader, never damage a PDF to manufacture a failure.
            failed_dir = root / 'failed-import'
            failed_dir.mkdir()
            (failed_dir / 'unavailable.pdf').write_bytes((root / 'papers' / 'check.pdf').read_bytes())
            with patch.object(service.store, 'import_pdf', side_effect=ValueError('controlled read failure')):
                service.store.sync_pdfs(failed_dir)
            unavailable = next(p for p in service.store.papers(project['id']) if p['title'] == 'unavailable')
            def check_missing_catalog(body):
                result = json.loads(body['messages'][-1]['content'])
                if 'materials' not in result:
                    result = json.loads(result['content'][0]['text'])
                missing = next(p for p in result['materials'] if p['id'] == unavailable['id'])
                assert missing['version_id'] is None and missing['status'].startswith('unavailable')
                return tool('read_material', {'paper_id':unavailable['id']})
            plans.extend([tool('set_scope',{'only_selected':False}), check_missing_catalog,
                          tool('read_material',{'paper_id':pdf_id}),
                          tool('write_file',{'title':'Partial materials','kind':'docx','content':f'Blue uses point clouds. [cite:{c["id"]}]\nOne material was unavailable.', 'citation_ids':[c['id']]}),
                          final('一篇不可用，其余正文支持成果。')])
            _, _, partial = submit('材料部分失败后继续', 'partial-materials')
            assert partial['files'] and any(t['name'] == 'read_material' and t['status'] == 'failed' for t in partial['tools'])
            previous_pdf_id = pdf_id
            pdf_id = unavailable['id']
            plans.extend([tool('set_scope',{'only_selected':True}), check_missing_catalog, final('所有选中材料均无可用正文，无法形成有依据的研究结论。')])
            _, _, empty = submit('所有材料不可用', 'all-unavailable')
            pdf_id = previous_pdf_id
            assert not empty['citations'] and not empty['files'] and not empty['coverage']
            assert unavailable['id'] not in {r['paper_id'] for r in service.store.reading_history(conversation)}
            plans.extend([tool('save_context', {'summary':'用户要求保留所有历史版本；研究进度保存在卡中。'}), final('已保存研究交接。')])
            submit('保留下一轮研究约定', 'handoff-memory')
            service.close()
            service = Service(root / 'papers', root / 'data')
            service.research.key = 'local-test-only'
            service.research.base_url = f'http://127.0.0.1:{provider.server_port}'
            handler.service = service
            assert binary(old_url) == original and binary(html_url) == html_bytes
            assert len(request(prefix)['versions']) == 2
            assert request(f'/api/projects/{project["id"]}/tasks/{planned_response["task_id"]}')['plan'] == complete_plan
            def recovered(body):
                user = json.loads(body['messages'][-1]['content'])
                assert user['recovered_dialogue'] and 'recovered_citations' not in user
                assert sum(len(m['text']) for m in user['recovered_dialogue']) <= 12000
                assert '保留所有历史版本' in user['handoff']['summary']
                return final('已从保存的对话及引用恢复。')
            plans.append(recovered)
            submit('继续刚才的内容', 'restart')
            def card_worker(body):
                context = json.loads(body['messages'][-1]['content'])
                assert set(context) == {'title','evidence'}
                assert len(context['evidence']) == 2
                assert not body.get('tools')
                assert body.get('thinking', {}).get('type') == 'disabled'
                assert all('paper_id' not in e and 'rect' not in e for e in context['evidence'])
                return final(json.dumps({'claims':[{'text':'Blue method uses point clouds.', 'quotes':[{'id':context['evidence'][0]['id'],'quote':context['evidence'][0]['quote']}]}], 'gaps':['No visual figure understanding.']}))
            def card_verifier(body):
                context = json.loads(body['messages'][-1]['content'])
                assert len(context['claims']) == 1 and not body.get('tools')
                assert body.get('thinking', {}).get('type') == 'disabled'
                assert 'point clouds' in context['evidence'][context['claims'][0]['evidence'][0]['id']]['quote']
                return final(json.dumps({'verdicts':[{'index':0,'supported':True,'reason':'Explicit source text.'}]}))
            def card_selector(body):
                context = json.loads(body['messages'][-1]['content'])
                assert set(context) == {'title','question','claims','gaps'} and not body.get('tools')
                assert context['claims'][0]['text'] == 'Blue method uses point clouds.'
                return final(json.dumps({'selected':[0],'gaps':context['gaps']}))
            plans.extend([tool('set_scope',{'only_selected':True}), tool('research_paper',{'paper_id':pdf_id,'question':'Compare method and evidence'}), card_worker, card_selector, card_verifier, final('研究卡已保存。')])
            _, _, card_task = submit('保存分篇研究卡', 'paper-card')
            assert card_task['coverage'][c['paper_version_id']]['read_blocks'] == 2
            assert service.store.one('SELECT complete FROM paper_cards WHERE conversation_id=?', (conversation,))['complete'] == 1
            service.close()
            service = Service(root / 'papers', root / 'data')
            service.research.key = 'local-test-only'
            service.research.base_url = f'http://127.0.0.1:{provider.server_port}'
            handler.service = service
            def check_reused(body):
                result = json.loads(body['messages'][-1]['content'])
                if 'reused' not in result:
                    result = json.loads(result['content'][0]['text'])
                assert result['reused'] and len(result['card']['claims']) == 1
                return tool('read_evidence', {'citation_ids':result['card']['claims'][0]['citation_ids']})
            plans.extend([tool('set_scope',{'only_selected':True}), tool('research_paper',{'paper_id':pdf_id,'question':'Compare method and evidence'}), check_reused, final('复用研究卡并回读依据。')])
            _, _, reused = submit('同一问题复用研究卡', 'reuse-paper-card')
            assert not reused['coverage']
            card_key = (conversation, c['paper_version_id'], 'Compare method and evidence')
            card_sql = 'SELECT * FROM paper_cards WHERE conversation_id=? AND paper_version_id=? AND question=?'
            original_card = json.loads(service.store.one(card_sql, card_key)['body'])
            correction_text = 'Evaluation is indoors only.'
            def revise_from_evidence(body):
                result = json.loads(body['messages'][-1]['content'])
                if 'evidence' not in result:
                    result = json.loads(result['content'][0]['text'])
                assert len(result['evidence']) == 1 and 'indoors only' in result['evidence'][0]['quote']
                return tool('revise_paper_card', {'paper_id':pdf_id, 'question':card_key[2], 'remove_texts':[],
                            'claims':[{'text':correction_text, 'citation_ids':[result['evidence'][0]['id']]}],
                            'gaps':['No visual figure understanding.']})
            def correction_verifier(body):
                context = json.loads(body['messages'][-1]['content'])
                assert not body.get('tools') and len(context['claims']) == 1
                assert context['claims'][0]['text'] == correction_text
                assert 'indoors only' in context['evidence'][context['claims'][0]['evidence'][0]['id']]['quote']
                return final(json.dumps({'verdicts':[{'index':0, 'supported':True, 'reason':'Explicit condition.'}]}))
            def correction_saved(body):
                result = json.loads(body['messages'][-1]['content'])
                if 'saved' not in result:
                    result = json.loads(result['content'][0]['text'])
                assert result['saved'] and len(result['card']['claims']) == 2
                assert result['card']['claims'][0]['text'] == original_card['claims'][0]['text']
                assert result['card']['claims'][1]['citation_ids'][0].startswith('E')
                return final('局部条件已修正，既有主张保留。')
            plans.extend([tool('set_scope',{'only_selected':True}), tool('read_material',{'paper_id':pdf_id,'query':'indoors'}),
                          revise_from_evidence, correction_verifier, correction_saved])
            _, _, revised_card_task = submit('只补充研究卡的评估条件', 'revise-paper-card')
            saved_card_row = dict(service.store.one(card_sql, card_key))
            revised_card = json.loads(saved_card_row['body'])
            assert revised_card['claims'][:-1] == original_card['claims']
            assert revised_card['claims'][-1]['text'] == correction_text
            assert saved_card_row['complete'] and saved_card_row['next_offset'] == 2
            assert not any(e['status'] == 'paper_read' for e in revised_card_task['events'])
            invalid_correction = {'paper_id':pdf_id, 'question':card_key[2],
                                  'remove_texts':[original_card['claims'][0]['text']],
                                  'claims':[{'text':'Evaluation includes outdoor trials.',
                                             'citation_ids':revised_card['claims'][-1]['citation_ids']}], 'gaps':[]}
            def rejected_correction_verifier(body):
                context = json.loads(body['messages'][-1]['content'])
                assert len(context['claims']) == 1 and not body.get('tools')
                assert context['claims'][0]['text'] == invalid_correction['claims'][0]['text']
                return final(json.dumps({'verdicts':[{'index':0, 'supported':False, 'reason':'Source says indoors only.'}]}))
            def correction_not_saved(body):
                result = json.loads(body['messages'][-1]['content'])
                if 'saved' not in result:
                    result = json.loads(result['content'][0]['text'])
                assert result['saved'] is False and result['verdicts'][0]['supported'] is False
                return final('不支持的修正未保存。')
            plans.extend([tool('set_scope',{'only_selected':True}), tool('revise_paper_card',invalid_correction),
                          rejected_correction_verifier, correction_not_saved])
            submit('失败的局部修正不能删除原主张或缺口', 'reject-paper-card-revision')
            assert dict(service.store.one(card_sql, card_key)) == saved_card_row
            print('PASS: local card correction verifies only additions, preserves prior claims and rejects changes atomically')
            # A new wording reuses full extraction; interrupted verification retains selection too.
            service.store.run('DELETE FROM claim_checks')
            plans.extend([tool('set_scope',{'only_selected':True}), tool('research_paper',{'paper_id':pdf_id,'question':'Resume verification'}), card_selector,
                          {'_http_error':{'status':402,'message':'Insufficient Balance'}}, final('核验失败，已保留提取和选择。')])
            submit('保存中间卡后发生核验故障', 'partial-paper-card')
            extraction = service.store.one('SELECT * FROM paper_extractions WHERE version_id=?', (c['paper_version_id'],))
            assert extraction['next_offset'] == 2 and extraction['complete']
            partial_card = service.store.one('SELECT * FROM paper_cards WHERE conversation_id=? AND question=?', (conversation,'Resume verification'))
            assert partial_card['next_offset'] == 2 and not partial_card['complete']
            plans.extend([tool('set_scope',{'only_selected':True}), tool('research_paper',{'paper_id':pdf_id,'question':'Resume verification'}), card_verifier, final('续接已完成。')])
            _, _, resumed_card = submit('从保存的卡续接核验', 'resume-paper-card')
            assert not resumed_card['coverage']
            assert service.store.one('SELECT complete FROM paper_cards WHERE conversation_id=? AND question=?', (conversation,'Resume verification'))['complete']
            candidate_pdf = pymupdf.open()
            candidate_page = candidate_pdf.new_page()
            for index in range(6):
                candidate_page.insert_text((60, 80 + index * 60), f'Evidence {index}: blue method uses point clouds indoors.')
            candidate_path = root / 'six-evidence.pdf'
            candidate_pdf.save(candidate_path)
            candidate_pdf.close()
            previous_pdf_id, pdf_id = pdf_id, service.store.import_pdf(project['id'], candidate_path)
            candidate_texts = [f'Candidate {index}: Blue method uses point clouds for indoor evaluation; performance is not established outdoors.' for index in range(75)]
            def many_candidates(body):
                context = json.loads(body['messages'][-1]['content'])
                assert len(context['evidence']) == 6 and not body.get('tools')
                ids = [e['id'] for e in context['evidence']]
                assert len(set(ids)) == 6
                candidate = {'claims':[{'text':text, 'quotes':[{'id':e['id'],'quote':e['quote']} for e in context['evidence'][:4]]} for text in candidate_texts],
                             'gaps':[f'Gap {index}' for index in range(13)]}
                return final(json.dumps(candidate))
            def select_candidates(body):
                context = json.loads(body['messages'][-1]['content'])
                assert [c['text'] for c in context['claims']] == candidate_texts
                assert len(context['gaps']) == 13 and not body.get('tools')
                return final(json.dumps({'selected':list(reversed(range(24))), 'gaps':[]}))
            verified_candidate_texts = []
            def verify_candidates(body):
                context = json.loads(body['messages'][-1]['content'])
                expected = list(reversed(candidate_texts[:24]))
                start = len(verified_candidate_texts)
                assert [c['text'] for c in context['claims']] == expected[start:start + 8]
                assert all(len({e['id'] for e in c['evidence']}) == 4 for c in context['claims'])
                verified_candidate_texts.extend(c['text'] for c in context['claims'])
                if len(verified_candidate_texts) < len(expected):
                    plans.appendleft(verify_candidates)
                return final(json.dumps({'verdicts':[{'index':c['index'], 'supported':c['text'] != expected[-1],
                                                      'reason':'Controlled unsupported selection.' if c['text'] == expected[-1] else 'Supported selection.'} for c in context['claims']]}))
            plans.extend([tool('set_scope',{'only_selected':True}), tool('research_paper',{'paper_id':pdf_id,'question':'Many valid intermediate candidates'}),
                          many_candidates, select_candidates, verify_candidates, final('只保存选中且核验通过的主张。')])
            submit('保留较多中间候选后选择并核验', 'many-paper-candidates')
            candidate_row = service.store.one('SELECT * FROM paper_cards WHERE conversation_id=? AND question=?', (conversation,'Many valid intermediate candidates'))
            candidate_card = json.loads(candidate_row['body'])
            assert candidate_row['complete'] and candidate_row['next_offset'] == 6
            assert [c['text'] for c in candidate_card['claims']] == verified_candidate_texts[:-1]
            assert len(candidate_card['claims']) == 23
            assert len(set(candidate_card['claims'][0]['citation_ids'])) == 4
            assert len(json_text(candidate_card)) <= 18000 and any(candidate_texts[0] in gap for gap in candidate_card['gaps'])
            large_card_revision = {'paper_id':pdf_id, 'question':'Many valid intermediate candidates', 'remove_texts':[],
                                   'claims':[{'text':correction_text, 'citation_ids':candidate_card['claims'][0]['citation_ids'][:1]}], 'gaps':[]}
            def large_card_verifier(body):
                context = json.loads(body['messages'][-1]['content'])
                assert len(context['claims']) == 1 and context['claims'][0]['text'] == correction_text
                assert 'indoors' in context['evidence'][context['claims'][0]['evidence'][0]['id']]['quote']
                return final(json.dumps({'verdicts':[{'index':0, 'supported':True, 'reason':'Explicit indoor condition.'}]}))
            plans.extend([tool('set_scope',{'only_selected':True}), tool('revise_paper_card',large_card_revision),
                          large_card_verifier, final('长度预算内的大卡局部修正已保存。')])
            submit('允许已完成研究卡局部修正', 'revise-many-paper-candidates')
            corrected = json.loads(service.store.one('SELECT body FROM paper_cards WHERE conversation_id=? AND question=?', (conversation,'Many valid intermediate candidates'))['body'])
            assert corrected['claims'][:-1] == candidate_card['claims'] and corrected['claims'][-1]['text'] == correction_text
            assert len(corrected['claims']) == 24 and len(json_text(corrected)) <= 18000
            pdf_id = previous_pdf_id
            print('PASS: 75 extracted candidates select 24, verify in batches, reject unsupported facts and retain valid corrections')
            # Exercise the bundled compactor using actual tool history and local provider usage.
            compaction_source = request(f'/api/projects/{project["id"]}/materials',
                                        {'title':'Compaction source', 'text':'OLD_COMPACTION_CONTEXT ' + 'source detail ' * 3000})
            previous_pdf_id, pdf_id = pdf_id, compaction_source['paper_id']
            compact_refs = []
            compact_requests = []
            def select_compaction_quote(body):
                result = json.loads(body['messages'][-1]['content'])
                if 'evidence' not in result:
                    result = json.loads(result['content'][0]['text'])
                # Keep the large read in history, but cite a genuine short selection.
                return tool('select_quote', {'citation_id':result['evidence'][0]['id'],
                                             'quote':'OLD_COMPACTION_CONTEXT source detail'})
            def create_pressure(body):
                result = json.loads(body['messages'][-1]['content'])
                if 'evidence' not in result:
                    result = json.loads(result['content'][0]['text'])
                assert 'OLD_COMPACTION_CONTEXT' in result['evidence'][0]['quote']
                compact_refs.append(result['evidence'][0]['id'])
                message = tool('list_paper_cards', {})
                message['content'] = 'RECENT_COMPACTION_TAIL ' + 'retain recent detail ' * 2000
                message['_usage'] = {'prompt_tokens':33000, 'completion_tokens':500, 'total_tokens':33500}
                return message
            def compact_summary(body):
                assert 'OLD_COMPACTION_CONTEXT' in json.dumps(body['messages'])
                compact_requests.append(body)
                message = final('COMPACTION_CHECKPOINT: scope remains selected; source reference ' + compact_refs[0])
                message['_compact'] = True
                message['_usage'] = {'prompt_tokens':16000, 'completion_tokens':30, 'total_tokens':16030}
                return message
            def after_compaction(body):
                messages = json.dumps(body['messages'])
                assert 'COMPACTION_CHECKPOINT' in messages and 'RECENT_COMPACTION_TAIL' in messages, [(m['role'], str(m.get('content',''))[:120], len(str(m.get('content','')))) for m in body['messages']]
                assert 'OLD_COMPACTION_CONTEXT' not in messages
                return tool('read_evidence', {'citation_ids':compact_refs})
            def finish_compaction(body):
                result = json.loads(body['messages'][-1]['content'])
                if 'evidence' not in result:
                    result = json.loads(result['content'][0]['text'])
                assert result['evidence'][0]['id'] == compact_refs[0]
                assert result['evidence'][0]['quote'] == 'OLD_COMPACTION_CONTEXT source detail'
                return final('压缩后仍可回读已保存依据。[cite:' + compact_refs[0] + ']')
            plans.extend([tool('set_scope',{'only_selected':True}), tool('read_material',{'paper_id':pdf_id}),
                          select_compaction_quote, create_pressure, compact_summary, after_compaction, finish_compaction])
            _, _, compacted = submit('验证长任务压缩及短引用恢复', 'automatic-compaction')
            pdf_id = previous_pdf_id
            assert len(compact_requests) == 1 and len(compacted['citations']) == 2
            precise = next(c for c in compacted['citations'] if c['quote'] == 'OLD_COMPACTION_CONTEXT source detail')
            assert precise['source_citation_id'] in {c['id'] for c in compacted['citations']}
            answer = service.store.one("SELECT text FROM messages WHERE task_id=? AND role='assistant'", (compacted['id'],))['text']
            assert f"[cite:{precise['id']}]" in answer
            compact_usage = [u for u in compacted['refs']['usage'] if u['role'] == 'compaction']
            assert len(compact_usage) == 1 and compact_usage[0]['calls'] == 1
            assert compact_usage[0]['inputTokens'] == 16000 and compact_usage[0]['outputTokens'] == 30
            print('PASS: actual automatic compaction replaces old tool history, preserves summary and citation resolution, and records usage')
            # #56: real Harness tools, HTTP persistence/downloads and independently opened files.
            from pptx import Presentation
            csv_source = request(f'/api/projects/{project["id"]}/materials',
                                 {'title':'Controlled chart fixture, not experimental results', 'text':'step,loss\n1,10.5\n2,7.25\n3,2.5\n'})
            csv_paper = request(f'/api/projects/{project["id"]}/papers/{csv_source["paper_id"]}')
            chart_spec = {'paper_id':csv_paper['id'], 'version_id':csv_paper['version_id'], 'x':'step', 'ys':['loss'], 'style':'line', 'claim':'Check supplied CSV values without interpolation'}
            chart_args = {'title':'Controlled data check', 'kind':'png', 'content':json.dumps(chart_spec), 'citation_ids':[]}
            chart_message_path = message_path
            plans.extend([tool('set_scope',{'only_selected':False}), tool('write_file',chart_args),
                          tool('write_file',{**chart_args,'content':json.dumps(chart_spec,indent=2)}), final()])
            chart_request, _, chart_task = submit('Generate a PNG from saved CSV', 'report-chart')
            assert len(chart_task['files']) == 1 and all(t['status'] == 'succeeded' for t in chart_task['tools'])
            chart_file = chart_task['files'][0]
            chart_ref = {k:chart_file[k] for k in ('artifact_id','version_id')}
            chart_base = f'/api/projects/{project["id"]}/artifacts/{chart_file["artifact_id"]}/versions/{chart_file["version_id"]}'
            chart_png = binary(chart_base + '/download')
            assert pymupdf.Pixmap(chart_png).width == 1600
            chart_sources = json.loads(binary(chart_base + '/sources'))
            assert chart_sources['report']['data']['series'][0]['values'] == [10.5,7.25,2.5]
            assert chart_sources['materials'][0]['paper_version_id'] == csv_paper['version_id']
            assert binary(chart_base + '/preview') == chart_png
            plans.extend([tool('set_scope',{'only_selected':False}),tool('write_file',{**chart_args,'content':json.dumps({**chart_spec,'style':'bar'})}),final()])
            _, _, mimir_chart = submit('Render a bar chart with Mimir', 'report-mimir-bar')
            assert len(mimir_chart['files']) == 1, mimir_chart['tools']
            bar_file = mimir_chart['files'][0]
            bar_base = f'/api/projects/{project["id"]}/artifacts/{bar_file["artifact_id"]}/versions/{bar_file["version_id"]}'
            bar_sources = json.loads(binary(bar_base + '/sources'))['report']
            assert bar_sources['renderer'].startswith('mimir-a568b373/metricFigureSvg')
            assert bar_sources['data']['series'][0]['values'] == [10.5,7.25,2.5]
            assert pymupdf.Pixmap(binary(bar_base + '/download')).width == 1600
            for suffix, values in [('negative',[-5,2]),('wide-number',[123456789012345,10000000000000])]:
                source = request(f'/api/projects/{project["id"]}/materials', {'title':suffix,'text':f'step,loss\ncontrol,{values[0]}\ntreatment,{values[1]}\n'})
                paper = request(f'/api/projects/{project["id"]}/papers/{source["paper_id"]}')
                spec = {**chart_spec,'paper_id':paper['id'],'version_id':paper['version_id'],'style':'bar'}
                plans.extend([tool('set_scope',{'only_selected':False}),tool('write_file',{**chart_args,'content':json.dumps(spec)}),final()])
                _, _, fallback = submit('Keep all bar labels readable',f'report-bar-{suffix}')
                assert len(fallback['files']) == 1, fallback['tools']
                file = fallback['files'][0]
                fallback_base = f'/api/projects/{project["id"]}/artifacts/{file["artifact_id"]}/versions/{file["version_id"]}'
                report = json.loads(binary(fallback_base + '/sources'))['report']
                assert report['renderer'].startswith('matplotlib-') and report['data']['series'][0]['values'] == values
                assert pymupdf.Pixmap(binary(fallback_base + '/download')).width == 1600
            assert request(message_path, chart_request)['duplicate']
            plans.extend([tool('set_scope',{'only_selected':True}), tool('write_file',chart_args), final('CSV is outside scope')])
            _, _, denied_chart = submit('Only selected paper', 'report-chart-scope')
            assert not denied_chart['files'] and any(t['status'] == 'failed' for t in denied_chart['tools'])
            invalid_csv = request(f'/api/projects/{project["id"]}/materials', {'title':'Missing value fixture', 'text':'step,loss\n1,NaN\n'})
            invalid_paper = request(f'/api/projects/{project["id"]}/papers/{invalid_csv["paper_id"]}')
            invalid_spec = {**chart_spec,'paper_id':invalid_paper['id'],'version_id':invalid_paper['version_id']}
            plans.extend([tool('set_scope',{'only_selected':False}), tool('write_file',{**chart_args,'content':json.dumps(invalid_spec)}), final('No finite data')])
            _, _, no_data = submit('Generate only from finite values', 'report-no-data')
            assert not no_data['files'] and any(t['status'] == 'failed' for t in no_data['tools'])
            for suffix, headers, accepted in [('mixed-units',['Speed [m/s]','Time [s]'],False), ('missing-units',['Speed','Time'],False), ('matching-units',['A [s]','B（s）'],True)]:
                unit_source = request(f'/api/projects/{project["id"]}/materials', {'title':suffix,'text':','.join(['step'] + headers) + '\n1,2,3\n'})
                unit_paper = request(f'/api/projects/{project["id"]}/papers/{unit_source["paper_id"]}')
                unit_spec = {**chart_spec,'paper_id':unit_paper['id'],'version_id':unit_paper['version_id'],'ys':headers}
                plans.extend([tool('set_scope',{'only_selected':False}),tool('write_file',{**chart_args,'content':json.dumps(unit_spec)}),final('Separate units')])
                _, _, unit_task = submit('Do not combine incompatible or unknown units',suffix)
                assert len(unit_task['files']) == int(accepted)
                assert any(t['status'] == 'failed' for t in unit_task['tools']) is not accepted
            actual_paper = request(f'/api/projects/{project["id"]}/papers/{pdf_id}')
            deck_spec = {'slides':[
                {'title':'真实数据文件检查', 'text':'These are controlled fixture values, not research results.', 'chart':chart_ref},
                {'title':'原始论文页', 'text':'Original page, not automatic figure identification.',
                 'figure':{'paper_id':pdf_id,'version_id':actual_paper['version_id'],'page':1}},
                {'title':'缺失资源', 'text':'Valid progress remains available.',
                 'figure':{'paper_id':pdf_id,'version_id':actual_paper['version_id'],'page':999}}]}
            deck_args = {'title':'组会文件检查', 'kind':'pptx','content':json.dumps(deck_spec),'citation_ids':[]}
            plans.extend([tool('set_scope',{'only_selected':False}), tool('write_file',deck_args), tool('write_file',deck_args), final()])
            _, _, deck_task = submit('Generate actual PPTX and preserve partial results', 'report-deck')
            assert len(deck_task['files']) == 1
            deck_file = deck_task['files'][0]
            deck_base = f'/api/projects/{project["id"]}/artifacts/{deck_file["artifact_id"]}/versions/{deck_file["version_id"]}'
            deck_bytes = binary(deck_base + '/download')
            deck = Presentation(io.BytesIO(deck_bytes))
            assert len(deck.slides) == 3
            charts = [s.chart for slide in deck.slides for s in slide.shapes if s.has_chart]
            assert len(charts) == 1 and list(charts[0].series[0].values) == [10.5,7.25,2.5]
            assert any(s.shape_type == 13 for s in deck.slides[1].shapes)
            sources = json.loads(binary(deck_base + '/sources'))
            assert len(sources['report']['warnings']) == 1
            committed_deck = json.loads(service.store.one('SELECT result FROM file_commits WHERE task_id=?', (deck_task['id'],))['result'])
            assert committed_deck['warnings'] == sources['report']['warnings'] and committed_deck['output_key']
            assert '图片未提取' in '\n'.join(s.text for s in deck.slides[2].shapes if s.has_text_frame)
            assert json.loads(deck.slides[0].notes_slide.notes_text_frame.text)['chart']['version_id'] == chart_file['version_id']
            rejected(f'/api/projects/{other}/artifacts/{deck_file["artifact_id"]}/versions/{deck_file["version_id"]}/download')
            rejected(f'/api/projects/{other}/artifacts/{deck_file["artifact_id"]}/versions/{deck_file["version_id"]}/sources')
            rejected(f'/api/projects/{other}/conversations/{request(f"/api/projects/{other}")["conversations"][0]["id"]}/messages',
                     {'text':'Wrong project','client_message_id':'report-cross','report_sources':[chart_ref]})
            plans.extend([tool('set_scope',{'only_selected':False}), tool('read_file',{'artifact_id':deck_file['artifact_id'],'version_id':deck_file['version_id']}),
                          tool('write_file',{**deck_args,'title':'Updated meeting','artifact_id':deck_file['artifact_id'],'base_version_id':deck_file['version_id']}), final()])
            _, _, revised_deck = submit('Revise meeting file', 'report-deck-v2')
            assert revised_deck['files'][0]['version_no'] == 2 and binary(deck_base + '/download') == deck_bytes
            # Reuse a chart from another conversation through project-wide reads, without a new chart.
            report_conversation = request(f'/api/projects/{project["id"]}/conversations',{})['conversation_id']
            message_path = f'/api/projects/{project["id"]}/conversations/{report_conversation}/messages'
            reused_spec = {'slides':[{'title':'已有成果汇报', 'text':'This uses an existing chart version.', 'chart':chart_ref, 'sources':[chart_ref]}]}
            plans.extend([tool('set_scope',{'only_selected':False}), tool('list_files',{'project_wide':True}),
                          tool('read_file',{**chart_ref,'project_wide':True}), tool('write_file',{**deck_args,'content':json.dumps(reused_spec)}), final()])
            _, _, reused = submit('Report from selected project artifact', 'report-existing', report_sources=[chart_ref])
            assert len(reused['files']) == 1 and reused['files'][0]['kind'] == 'pptx'
            assert reused['refs']['report_sources'] == [chart_ref]
            plain_spec = {'slides':[{'title':'已有笔记与进展','text':'Experiment results are not yet available. Next: collect the measurements.'}]}
            long_title = '组会汇报' * 30
            plans.extend([tool('write_file',{**deck_args,'title':long_title,'content':json.dumps(plain_spec)}),final()])
            _, _, plain = submit('PPT from existing notes without figures', 'report-independent')
            assert len(plain['files']) == 1
            plain_file = plain['files'][0]
            plain_deck = Presentation(io.BytesIO(binary(f'/api/projects/{project["id"]}/artifacts/{plain_file["artifact_id"]}/versions/{plain_file["version_id"]}/download')))
            assert plain_deck.core_properties.title == long_title
            assert any(shape.text.endswith('…') for shape in plain_deck.slides[0].shapes if shape.has_text_frame)
            foreign = request(f'/api/projects/{other}/materials', {'title':'Other project input','text':'Private input'})
            foreign_paper = request(f'/api/projects/{other}/papers/{foreign["paper_id"]}')
            foreign_spec = {'slides':[{'title':'Forbidden input','text':'Do not include it','figure':{'paper_id':foreign_paper['id'],'version_id':foreign_paper['version_id'],'page':1}}]}
            plans.extend([tool('set_scope',{'only_selected':False}),tool('write_file',{**deck_args,'content':json.dumps(foreign_spec)}),final('Rejected')])
            _, _, foreign_task = submit('Reject another project even for optional images', 'report-foreign')
            assert not foreign_task['files'] and any(t['status'] == 'failed' for t in foreign_task['tools'])
            oversized = {'slides':[{'title':'Long caption','text':'Long caption must not silently overflow. ' * 150, 'chart':chart_ref}]}
            plans.extend([tool('set_scope',{'only_selected':False}),tool('write_file',{**deck_args,'content':json.dumps(oversized)}),final('Shorten the slide')])
            _, _, too_long = submit('Reject overflowing captions', 'report-overflow')
            assert not too_long['files']
            service.close()
            service = Service(root / 'papers', root / 'data')
            service.research.key = 'local-test-only'
            service.research.base_url = f'http://127.0.0.1:{provider.server_port}'
            handler.service = service
            assert binary(deck_base + '/download') == deck_bytes and binary(chart_base + '/download') == chart_png
            assert request(chart_message_path, chart_request)['duplicate']
            print('PASS: #56 real PNG/PPTX files, source data, editable chart values, partial failures, immutable revisions, project-wide sources and cross-project rejection')
            message_path = chart_message_path
            original_output = {'title':'Resume deliverable','kind':'docx','content':'Original committed wording.','citation_ids':[]}
            plans.extend([tool('set_scope',{'only_selected':True}),tool('write_file',original_output),
                          {'_http_error':{'status':402,'message':'Controlled failure after commit'}}])
            _, _, interrupted_output = submit('生成一份恢复检查文件','explore-resume',expected='failed')
            assert len(interrupted_output['files']) == 1
            plans.extend([tool('set_scope',{'only_selected':True}),tool('write_file',{**original_output,'content':'Regenerated different wording.'}),final()])
            request(f'/api/projects/{project["id"]}/conversations/{conversation}/tasks/{interrupted_output["id"]}/resume',{})
            deadline = time.monotonic() + 90
            while time.monotonic() < deadline:
                resumed = request(f'/api/projects/{project["id"]}/tasks/{interrupted_output["id"]}')
                if resumed['status'] in ('succeeded','failed'):
                    break
                time.sleep(.1)
            assert resumed['status'] == 'succeeded' and len(resumed['files']) == 1, resumed.get('error')
            assert len([e for e in resumed['events'] if e['status'] == 'committed']) == 1
            saved_output = request(f'/api/projects/{project["id"]}/artifacts/{resumed["artifact_id"]}')
            assert saved_output['versions'][0]['body'] == original_output['content']
            print('PASS: actual exploratory Harness resume keeps the logical output despite changed generated wording')
            # Same public message entry, actual SDK; only model HTTP is controlled.
            for skill, heading in [('research-lit-review','# 文献综述'),
                                   ('research-experiment-plan','# 实验设计'),
                                   ('research-result-to-claim','# 结果分析')]:
                def skill_response(body):
                    sent = json.dumps(body, ensure_ascii=False)
                    assert heading in sent and 'MethodAtlas 内置研究成果规范' in sent, 'new built-in Skill not loaded'
                    context = json.loads(body['messages'][-1]['content'])
                    cid = context['evidence'][0]['evidence'][0]['id']
                    return final(json.dumps({'answer':'已根据原文生成。', 'files':[{
                        'title':skill, 'kind':'docx', 'content':f'原文依据。[cite:{cid}]', 'citation_ids':[cid]}]}))
                def skill_review(body):
                    context = json.loads(body['messages'][-1]['content'])
                    assert context['mode'] == 'document' and context['draft']['title'] == skill
                    qid, quote = next(iter(context['quotes'].items()))
                    return final(json.dumps({'content':f'{quote["text"]}[cite:{qid}]', 'citation_ids':[qid]}))
                plans.extend([{**final(json.dumps({'intent':'research','mode':'rag','only_selected':True,
                    'skills':[skill], 'reads':[{'paper_id':pdf_id,'query':'evidence'}], 'files':[]})), '_rag':True}, skill_response, skill_review])
                sent, _, outcome = submit('根据材料生成研究成果：' + skill, skill)
                assert len(outcome['files']) == 1
                assert any(e['status'] == 'skill_loaded' for e in outcome['events'])
                assert request(message_path, sent)['duplicate']
                version = request(f'/api/projects/{project["id"]}/artifacts/{outcome["artifact_id"]}')['versions'][0]
                assert version['citations'] and version['materials']
                assert version['citations'][0]['quote'] in version['body']
                assert any(e['status'] == 'review' for e in outcome['events'])
                file = outcome['files'][0]
                assert binary(f'/api/projects/{project["id"]}/artifacts/{file["artifact_id"]}/versions/{file["version_id"]}/download').startswith(b'PK')
            print('PASS: three first-party built-in Skills load through actual Harness, with versioned cited files and duplicate protection')
            plans.extend([{**final(json.dumps({'intent':'chat','mode':'rag','only_selected':True,
                'skills':['research-experiment-plan'], 'reads':[], 'files':[]})), '_rag':True},
                final(json.dumps({'answer':'只讨论：先确定主张和评价指标。','files':[
                    {'title':'Unrequested','kind':'docx','content':'Must not save','citation_ids':[]}]}))])
            _, _, discussion = submit('只讨论实验设计，不生成文件', 'skill-discussion')
            assert not discussion['files'] and discussion['kind'] == 'chat'
            def explore_skill(body):
                assert '# 结果分析' in json.dumps(body, ensure_ascii=False)
                return final('缺少真实实验结果，尚不能判断主张；未执行实验。')
            plans.extend([{**final(json.dumps({'intent':'research','mode':'explore','only_selected':True,
                'skills':['research-result-to-claim'],'reason':'根据中途发现核对结果'})), '_rag':True}, explore_skill])
            _, _, missing_results = submit('分析结果，缺少实验记录时说明', 'skill-missing-results')
            assert not missing_results['files'] and any(e['status'] == 'skill_loaded' for e in missing_results['events'])
            plans.append({**final(json.dumps({'intent':'research','mode':'rag','only_selected':True,
                'skills':['../../unsupported'],'reads':[], 'files':[]})), '_rag':True})
            _, _, unknown_skill = submit('非法方法检查', 'skill-invalid', expected='failed')
            assert unknown_skill['error_info']['code'] == 'input' and not unknown_skill['files']
            print('PASS: Skill discussions cannot write, exploration loads the same method, missing results stay explicit and unknown Skills fail closed')
            plans.append({'_http_error':{'status':402,'message':'Insufficient Balance'}})
            _, _, no_balance = submit('余额故障检查', 'no-balance', expected='failed')
            assert no_balance['error_info']['code'] == 'billing' and 'HTTP' not in no_balance['error'] and '余额不足' in no_balance['error'] and not no_balance['files']
            assert not plans
            print('PASS: actual Harness + local provider, project tools, scope enforcement, answer without files, immutable PDF bytes, precise citations, ownership, duplicate requests/tools, revisions, rejected stale writes, HTML isolation and restart recovery')
        finally:
            server.shutdown()
            server.server_close()
            service.close()
            provider.shutdown()
            provider.server_close()


if __name__ == '__main__':
    main()
