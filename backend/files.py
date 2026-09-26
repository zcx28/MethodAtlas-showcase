"""Editable DOCX and isolated HTML, promoted only after successful validation."""
import html
import io
import re
import secrets
from html.parser import HTMLParser

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from .state import canonical_citation_id
from .methods import STYLE


class SafeHTML(HTMLParser):
    tags = set('html head body title main article section header footer nav div span p h1 h2 h3 h4 h5 h6 ul ol li strong em b i br hr pre code blockquote table thead tbody tr th td caption a button details summary style svg g path rect circle ellipse line polyline polygon text tspan defs marker lineargradient radialgradient stop'.split())
    attrs = set('id class title scope role aria-label aria-hidden aria-expanded hidden data-label data-method data-related data-detail data-close-detail colspan rowspan type open data-citation viewbox xmlns x y x1 x2 y1 y2 width height d r rx ry cx cy points fill stroke stroke-width opacity fill-opacity stroke-opacity transform text-anchor font-size font-family font-weight dominant-baseline offset stop-color marker-end marker-start markerwidth markerheight refx refy orient preserveaspectratio'.split())

    def __init__(self, citation_ids, flat=False):
        super().__init__(convert_charrefs=True)
        self.parts, self.blocked, self.citation_ids = [], 0, set(citation_ids)
        self.in_style = False
        self.flat = flat

    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'iframe', 'object', 'embed', 'template') or self.flat and tag == 'style':
            self.blocked += 1
        if self.blocked or tag not in self.tags:
            return
        self.in_style = tag == 'style'
        tag = {'details':'section', 'summary':'h2'}.get(tag, tag)
        safe = []
        is_citation = any(key == 'data-citation' for key, _ in attrs)
        if self.flat and tag in ('button','nav') and not is_citation:
            raise ValueError('内置成果只接受正文和引用；请移除操作按钮与导航')
        for key, value in attrs:
            value = value or ''
            if key in ('hidden','open','aria-expanded') or self.flat and key == 'style':
                continue
            if key == 'data-citation':
                value = canonical_citation_id(value)
                if value not in self.citation_ids:
                    raise ValueError('HTML 含有未读取或不属于当前对话的引用')
            if is_citation and key == 'role':
                continue
            if key in self.attrs or key == 'style' or key == 'href' and value.startswith('#'):
                # SVG attributes are case-sensitive in serialized markup.
                key = {'viewbox': 'viewBox', 'markerwidth': 'markerWidth', 'markerheight': 'markerHeight', 'refx': 'refX', 'refy': 'refY', 'preserveaspectratio': 'preserveAspectRatio'}.get(key, key)
                safe.append(f'{key}="{html.escape(value, quote=True)}"')
        for key, value in attrs:
            if key in ('data-method','data-detail','data-close-detail') and not re.fullmatch(r'\d+', value or '') or key == 'data-related' and not re.fullmatch(r'p\d+( p\d+)*', value or ''):
                raise ValueError('研究交互标识无效')
        if is_citation:
            safe.append('tabindex="0"')
            if tag not in ('a','button'):
                safe.append('role="button"')
        self.parts.append('<' + tag + (' ' + ' '.join(safe) if safe else '') + '>')

    def handle_endtag(self, tag):
        if tag == 'style':
            self.in_style = False
        if tag in ('script', 'iframe', 'object', 'embed', 'template') or self.flat and tag == 'style':
            self.blocked = max(0, self.blocked - 1)
        elif not self.blocked and tag in self.tags:
            self.parts.append('</' + {'details':'section', 'summary':'h2'}.get(tag, tag) + '>')

    def handle_data(self, data):
        if not self.blocked:
            self.parts.append(data if self.in_style else html.escape(data))


def render_file(kind, title, content, citations, *, flat=False):
    if not isinstance(content, str) or not content.strip() or len(content) > 500000:
        raise ValueError('文件内容无效（最多 500000 字符）')
    sources = {}
    for index, citation in enumerate(citations, 1):
        sources.setdefault(citation['paper_version_id'], []).append(index)
    refs = [f"[{', '.join(map(str, indices))}] {' '.join(citations[indices[0]-1]['title'].split())}" for indices in sources.values()]
    if kind == 'html':
        parser = SafeHTML((c['id'] for c in citations), flat=flat)
        parser.feed(content)
        nonce = secrets.token_hex(16)
        policy = f"default-src 'none'; style-src 'unsafe-inline'; img-src data:; script-src 'nonce-{nonce}'; connect-src 'none'; base-uri 'none'; form-action 'none'"
        bridge = "function cite(e){const n=e.target.closest('[data-citation]');if(!n)return;e.preventDefault();if(parent!==window){parent.postMessage({type:'methodatlas-citation',id:n.dataset.citation},'*')}else{const ref=document.getElementById('ref-'+n.dataset.citation);if(ref){ref.scrollIntoView({block:'center'});ref.focus({preventScroll:true})}}}document.addEventListener('click',cite);document.addEventListener('keydown',e=>{if((e.key==='Enter'||e.key===' ')&&!['A','BUTTON'].includes(e.target.tagName))cite(e)});"
        references = ''.join(f'<section tabindex="-1" id="ref-{html.escape(c["id"])}"><h3>[{i}] {html.escape(c["title"])} · 第 {c.get("page", "未知")} 页</h3><blockquote>{html.escape(c.get("quote", "未保存原文片段"))}</blockquote></section>' for i,c in enumerate(citations,1))
        appendix = f'<section class="methodatlas-sources"><h2>参考文献</h2>{references}</section>' if references else ''
        markup = ''.join(parser.parts)
        if not re.search(r'<h1(?:\s|>)', markup, re.I):
            markup = f'<h1>{html.escape(title)}</h1>' + markup
        body = f'<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta http-equiv="Content-Security-Policy" content="{html.escape(policy, quote=True)}"><title>{html.escape(title)}</title></head><body class="reading-view">{markup}{appendix}<style>{STYLE}</style><script nonce="{nonce}">{bridge}</script></body></html>'
        return body.encode(), body
    if kind != 'docx':
        raise ValueError('仅支持 docx 或 html 文件')
    # PDF extraction can contain XML-forbidden controls. Clean the export, not evidence.
    title, content, *refs = [re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]', '', text)
                             for text in (title, content, *refs)]
    if re.search(r'\[\d+\]', content):
        raise ValueError('DOCX 引用请使用 [cite:引用id]，不要手写 [n]；编号由系统生成')
    numbers = {c['id']: index for index, c in enumerate(citations, 1)}
    def number_reference(match):
        citation_id = canonical_citation_id(match[1])
        if citation_id not in numbers:
            raise ValueError('DOCX 含有未列入 citation_ids 的引用')
        return f'[{numbers[citation_id]}]'
    content = re.sub(r'\[cite:([^\]]+)\]', number_reference, content)
    if '[cite:' in content:
        raise ValueError('DOCX 引用标记格式错误')
    document = Document()
    document.sections[0].page_width, document.sections[0].page_height = Inches(8.27), Inches(11.69)
    for section in document.sections:
        section.top_margin = section.bottom_margin = Inches(.8)
        section.left_margin = section.right_margin = Inches(.8)
    document.styles['Normal'].font.size = Pt(11)
    document.styles['Normal'].paragraph_format.line_spacing = 1.5
    document.styles['Normal'].paragraph_format.space_after = Pt(8)
    for style_name in ('Title','Heading 1','Heading 2','Heading 3'):
        style = document.styles[style_name]
        style.font.color.rgb = RGBColor(0,0,0)
        style.font.underline = False
        for border in style.element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
    document.add_heading(title, 0)
    if flat:
        content = re.sub(r'\A\s*#\s+[^\n]*(?:\n|$)', '', content, count=1)
    lines = content.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if line.startswith('|') and index + 1 < len(lines) and re.match(r'^\|[\s:|\-]+\|$', lines[index + 1].strip()):
            rows = []
            while index < len(lines) and lines[index].strip().startswith('|'):
                row = lines[index].strip()
                if not re.match(r'^\|[\s:|\-]+\|$', row):
                    rows.append([cell.strip() for cell in row.strip('|').split('|')])
                index += 1
            if not rows or len(rows[0]) > 12 or any(len(row) != len(rows[0]) for row in rows):
                raise ValueError('DOCX 表格列数不一致或超过 12 列')
            table = document.add_table(rows=0, cols=len(rows[0]))
            table.style = 'Light Shading Accent 1'
            for row in rows:
                for cell, value in zip(table.add_row().cells, row):
                    cell.text = value
            continue
        heading = re.match(r'^(#{1,6})\s+(.+)', line)
        if heading and heading[1] == '#' and heading[2].strip() == title.strip():
            index += 1
            continue
        if heading:
            document.add_heading(heading[2], min(3, len(heading[1])))
        elif line:
            document.add_paragraph(line[2:] if line.startswith('- ') else line,
                                   style='List Bullet' if line.startswith('- ') else None)
        index += 1
    if refs:
        document.add_heading('原文依据', 1)
        for ref in refs:
            document.add_paragraph(ref)
    buffer = io.BytesIO()
    document.save(buffer)
    raw = buffer.getvalue()
    Document(io.BytesIO(raw))
    return raw, content
