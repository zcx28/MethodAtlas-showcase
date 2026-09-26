"""Public document boundary: immutable versions, review conflicts and exports."""
import tempfile
import copy
import base64
import io
import json
import threading
import os
import sqlite3
import uuid
from collections import deque
from contextlib import closing
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from docx import Document
import pymupdf
from .state import Store
from .writing import Writing, selection_slice
from .agent import Research


def rejects(call):
    try:
        call()
    except ValueError:
        return
    raise AssertionError('Invalid or conflicting operation was accepted')


def main():
    selected = selection_slice([{'id':'a','type':'p','children':[{'text':'😀abc𠮷'}]}], {'anchor':{'path':[0,0],'offset':2},'focus':{'path':[0,0],'offset':3}})
    assert selected[1] == 'a' and selected[2][0]['text'] == '😀' and selected[3][0]['text'] == 'bc𠮷'
    rejects(lambda: selection_slice([{'id':'a','type':'p','children':[{'text':'😀abc'}]}], {'anchor':{'path':[0,0],'offset':1},'focus':{'path':[0,0],'offset':3}}))
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
        store = Store(Path(directory) / 'state.sqlite3')
        project = store.create_project('写作检查')
        writing = Writing(store)
        backup = next((Path(directory) / 'backups').glob('before-writing-*.sqlite3'))
        with closing(sqlite3.connect(backup)) as source, closing(sqlite3.connect(Path(directory) / 'restored.sqlite3')) as restored:
            source.backup(restored)
            assert restored.execute('PRAGMA integrity_check').fetchone()[0] == 'ok'
            assert restored.execute('SELECT name FROM projects WHERE id=?', (project,)).fetchone()[0] == '写作检查'
        value = [{'id': 'intro', 'type': 'p', 'children': [{'text': '中文 42。'}]}]
        # A newly imported paper can be registered and saved without an AI run.
        paper_id = store.paste(project, '新增引用论文', '可追溯原文')
        paper = store.paper(project, paper_id)
        ref_doc = writing.save(project, None, {'request_id':'reference-doc','title':'引用检查','document':value})
        body = {'paper_id':paper_id, 'paper_version_id':paper['version_id']}
        citation = writing.reference(project, ref_doc['artifact_id'], body)
        assert writing.reference(project, ref_doc['artifact_id'], body)['id'] == citation['id']
        assert citation['kind'] == 'bibliography' and citation['quote'] == ''
        bad_project = store.create_project('引用隔离')
        foreign = store.paste(bad_project, '其他项目', '其他原文')
        rejects(lambda: writing.reference(project, ref_doc['artifact_id'], {'paper_id':foreign,'paper_version_id':store.paper(bad_project,foreign)['version_id']}))
        cited = copy.deepcopy(value)
        cited[0]['children'].append({'type':'citation','citation_id':citation['id'],'children':[{'text':''}]})
        saved_ref = writing.save(project, ref_doc['artifact_id'], {'request_id':'reference-save','base_version_id':ref_doc['version_id'],'title':'引用检查','document':cited})
        assert writing.version(project, ref_doc['artifact_id'])['citations'][0]['paper_version_id'] == paper['version_id']
        assert citation['id'].encode() in writing.export(project, ref_doc['artifact_id'], saved_ref['version_id'], 'html')[0]
        first = writing.save(project, None, {'request_id': 'create-1', 'title': '论文', 'document': value})
        assert writing.save(project, None, {'request_id': 'create-1', 'title': '论文', 'document': value}) == first
        version = store.artifact(project, first['artifact_id'])['versions'][0]
        assert version['payload']['document'] == value
        original = writing.export(project, first['artifact_id'], first['version_id'], 'html')[0]
        assert original.startswith(b'<!doctype html>')
        other = store.create_project('隔离项目')
        rejects(lambda: writing.get(other, first['artifact_id']))
        rejects(lambda: writing.export(other, first['artifact_id'], first['version_id'], 'docx'))
        aid = first['artifact_id']
        values = value + [{'id':'tail','type':'p','children':[{'text':'未受影响的段落'}]}]
        second = writing.save(project, aid, {'request_id':'manual-2','base_version_id':first['version_id'],'title':'论文','document':values})
        rejects(lambda: writing.save(project, aid, {'request_id':'stale','base_version_id':first['version_id'],'title':'论文','document':value}))
        assert writing.export(project, aid, first['version_id'], 'html')[0] == original

        plans, calls = deque(), []
        class Provider(BaseHTTPRequestHandler):
            def log_message(self, *_):
                pass
            def do_POST(self):
                body = json.loads(self.rfile.read(int(self.headers['Content-Length'])))
                calls.append(body)
                assert not body.get('tools'), 'Writing must not acquire shell or network tools'
                context = json.loads(next(m['content'] for m in reversed(body['messages']) if m['role'] == 'user'))
                if 'known_reason' in context:
                    response = {'explanation':'这次没有取得完整内容，请缩小修改范围后发送新请求。'}
                else:
                    assert 'Paper De-AI' in json.dumps(body, ensure_ascii=False), 'Original Mimir skill was not loaded'
                    response = plans.popleft()
                chunk = {'id':'writing-check','object':'chat.completion.chunk','created':0,'model':body['model'],'choices':[{'index':0,'delta':{'role':'assistant','content':json.dumps(response,ensure_ascii=False)},'finish_reason':None}]}
                done = {**chunk, 'choices':[{'index':0,'delta':{},'finish_reason':'stop'}], 'usage':{'prompt_tokens':100,'completion_tokens':20,'total_tokens':120}}
                raw = ('data: '+json.dumps(chunk)+'\n\ndata: '+json.dumps(done)+'\n\ndata: [DONE]\n\n').encode()
                self.send_response(200); self.send_header('Content-Type','text/event-stream'); self.send_header('Content-Length',str(len(raw))); self.end_headers(); self.wfile.write(raw)
        provider = ThreadingHTTPServer(('127.0.0.1',0), Provider)
        threading.Thread(target=provider.serve_forever, daemon=True).start()
        research = Research(store)
        research.key, research.base_url = 'local-check', f'http://127.0.0.1:{provider.server_port}'
        writing = Writing(store, research)
        try:
            selection = {'anchor':{'path':[0,0],'offset':0},'focus':{'path':[0,0],'offset':2}}
            request = {'request_id':'proposal-1','base_version_id':second['version_id'],'instruction':'只润色选中文字','mode':'polish','selection':selection}
            plans.append({'text':'表述','summary':'精简表述'})
            reviewed = writing.propose(project, aid, request)
            assert reviewed['proposals'][-1]['status'] == 'pending', reviewed['proposals'][-1]
            assert len(reviewed['versions']) == 2
            research.close(); store.close()
            store = Store(Path(directory) / 'state.sqlite3')
            research = Research(store)
            research.key, research.base_url = 'local-check', f'http://127.0.0.1:{provider.server_port}'
            writing = Writing(store, research)
            assert writing.get(project, aid)['proposals'][-1]['status'] == 'pending'
            writing.propose(project, aid, request)
            assert len(calls) == 1
            rejects(lambda: writing.export(project, aid, first['version_id'], 'pdf'))
            changed = copy.deepcopy(values); changed[1]['children'][0]['text'] = '人工内容要保留'
            writing.save(project, aid, {'request_id':'manual-3','base_version_id':second['version_id'],'title':'论文','document':changed})
            accepted = writing.decide(project, aid, 'proposal-1', 'accept')
            assert accepted['versions'][-1]['payload']['document'][0]['children'][1]['text'] == '表述'
            assert accepted['versions'][-1]['payload']['document'][1]['children'][0]['text'] == '人工内容要保留'
            assert len(writing.decide(project, aid, 'proposal-1', 'accept')['versions']) == 4
            assert writing.export(project, aid, first['version_id'], 'html')[0] == original

            current = accepted['versions'][-1]
            replacement = copy.deepcopy(current['payload']['document']); replacement[0]['children'] = [{'text':'正文 42。'}]
            plans.append({'document':replacement,'summary':'修改首段'})
            request = {'request_id':'proposal-2','base_version_id':current['id'],'instruction':'润色首段','mode':'polish'}
            assert writing.propose(project, aid, request)['proposals'][-1]['status'] == 'pending'
            manual = copy.deepcopy(current['payload']['document']); manual[0]['children'] = [{'text':'同段人工改动 42。'}]
            writing.save(project, aid, {'request_id':'manual-4','base_version_id':current['id'],'title':'论文','document':manual})
            rejects(lambda: writing.decide(project, aid, 'proposal-2', 'accept'))
            count = len(writing.get(project, aid)['versions'])
            writing.decide(project, aid, 'proposal-2', 'reject')
            assert len(writing.decide(project, aid, 'proposal-2', 'reject')['versions']) == count
            current = writing.get(project, aid)['versions'][-1]
            plans.append({'document':value,'summary':'错误地改动数字'})
            bad = copy.deepcopy(value); bad[0]['children'][0]['text'] = '中文 99。'; plans[-1]['document'] = bad
            request = {'request_id':'proposal-bad','base_version_id':current['id'],'instruction':'润色','mode':'polish'}
            assert writing.propose(project, aid, request)['proposals'][-1]['status'] == 'failed'
            assert len(writing.get(project, aid)['versions']) == count

            png = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0,0,12,12)); png.clear_with(150)
            structure = [
                {'id':'h','type':'h1','children':[{'text':'研究方法'}]},
                {'id':'p','type':'p','children':[{'text':'中文与 English 42','bold':True}]},
                {'id':'list','type':'ol','children':[{'type':'li','children':[{'text':'步骤一'}]}]},
                {'id':'table','type':'table','children':[{'type':'tr','children':[{'type':'td','children':[{'type':'p','children':[{'text':'42'}]}]}]}]},
                {'id':'img','type':'img','url':'data:image/png;base64,'+base64.b64encode(png.tobytes('png')).decode(),'alt':'验证图','children':[{'text':''}]},
                {'id':'eq','type':'equation','formula':'E=mc^2','children':[{'text':''}]},
            ]
            saved = writing.save(project, aid, {'request_id':'structures','base_version_id':current['id'],'title':'论文结构','document':structure})
            docx = Document(io.BytesIO(writing.export(project, aid, saved['version_id'], 'docx')[0]))
            assert docx.tables[0].cell(0,0).text == '42' and len(docx.inline_shapes) == 2
            assert any('研究方法' in p.text for p in docx.paragraphs)
            raw = writing.export(project, aid, saved['version_id'], 'pdf')[0]
            with pymupdf.open(stream=raw) as pdf:
                assert '42' in ''.join(p.get_text() for p in pdf) and len(pdf) > 0
            restored = writing.restore(project, aid, {'request_id':'restore','version_id':first['version_id'],'base_version_id':saved['version_id']})
            assert restored['version_no'] == saved['version_no'] + 1
            assert writing.get(project, aid)['versions'][-1]['payload']['document'] == value
            assert store.messages(store.conversations(project)[0]['id']) == [], 'Right-side edits leaked into chat'
        finally:
            research.close(); provider.shutdown(); provider.server_close()
        store.close()
        store = Store(Path(directory) / 'state.sqlite3')
        assert Writing(store).get(project, aid)['proposals'][0]['status'] == 'accepted'
        store.close()
    print('Writing checks passed')


def real_check(source_db):
    if not os.getenv('DEEPSEEK_API_KEY'):
        raise SystemExit('DEEPSEEK_API_KEY missing; no model call made')
    root = Path('.check-data') / ('writing-real-' + uuid.uuid4().hex[:8])
    root.mkdir(parents=True)
    store = Store(root / 'state.sqlite3')
    project = store.create_project('论文写作真实验证')
    # Copy an already parsed immutable paper version, not the source project or
    # conversations. The source database is opened read-only and never recovered.
    source = sqlite3.connect(Path(source_db).resolve().as_uri() + '?mode=ro', uri=True)
    source.row_factory = sqlite3.Row
    try:
        paper = dict(source.execute("SELECT * FROM papers WHERE title LIKE '04_Diffusion_Policy%' LIMIT 1").fetchone())
        version = dict(source.execute('SELECT * FROM paper_versions WHERE id=?', (paper['current_version_id'],)).fetchone())
        paper['project_id'] = project
        with store.transaction() as db:
            for table, row in [('papers',paper),('paper_versions',version)]:
                db.execute(f'INSERT INTO {table} ({",".join(row)}) VALUES ({",".join("?" for _ in row)})', tuple(row.values()))
    finally:
        source.close()
    research = Research(store); research.model = 'deepseek-v4-flash'
    writing = Writing(store, research)
    try:
        saved = writing.save(project,None,{'request_id':'real-create','title':'Diffusion Policy 方法说明','document':[{'id':'intro','type':'p','children':[{'text':'待起草。'}]}]})
        aid = saved['artifact_id']
        reports = []
        for index, (mode, instruction) in enumerate([
            ('draft', '仅基于 Diffusion Policy visuomotor policy 已提供原文起草约100字的方法说明。不得宣称做过实验；关键结论插入真实引用。保留缺口。'),
            ('polish', '润色成简洁中文学术表述，保留全部事实、数字和引用，不新增信息。'),
            ('rebuttal', '审稿意见：请报告新的消融实验结果。目前我们没有做过任何新实验。写一段诚实的审稿回复，标明“需补实验”，不要编造结果或承诺执行。原有说明和引用保留。'),
        ]):
            current = writing.version(project, aid)
            request_id = 'real-' + str(index)
            item = writing.propose(project,aid,{'request_id':request_id,'base_version_id':current['id'],'instruction':instruction,'mode':mode,'only_selected':True,'selected_paper_ids':[paper['id']]})
            proposal = item['proposals'][-1]
            reports.append({k:proposal[k] for k in ('id','status','result','error','usage')})
            (root / 'usage.json').write_text(json.dumps(reports,ensure_ascii=False,indent=2),encoding='utf-8')
            print(mode, proposal['status'], proposal['error'] or '', flush=True)
            assert proposal['status'] == 'pending', proposal['error']
            assert len(item['versions']) == current['version_no']
            item = writing.decide(project,aid,request_id,'accept')
            selected = item['versions'][-1]
            for kind in ('html','pdf','docx'):
                raw, _ = writing.export(project,aid,selected['id'],kind)
                (root / f'{mode}-v{selected["version_no"]}.{kind}').write_bytes(raw)
            assert selected['citations'], 'Real source citations missing'
            for citation in selected['citations']:
                original_page = json.loads(version['pages'])[citation['page'] - 1]
                assert original_page['text'][citation['start']:citation['end']] == citation['quote']
                assert citation['paper_version_id'] == version['id']
        print(json.dumps({'data_dir':str(root.resolve()),'project_id':project,'artifact_id':aid,'runs':reports},ensure_ascii=False),flush=True)
    finally:
        research.close(); store.close()


if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument('--source-db', help='Run real DeepSeek writing against an existing Diffusion Policy version')
    args = parser.parse_args()
    real_check(args.source_db) if args.source_db else main()
