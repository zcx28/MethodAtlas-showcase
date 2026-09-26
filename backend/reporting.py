"""Deterministic research files from authorized, versioned inputs.

Calls Mimir's pinned metricFigureSvg and renderDeck modules through a JSON
bridge; originals and attribution live in third_party/mimir. No model code runs.
"""
from .errors import failure_message, user_message
import csv
import hashlib
import io
import json
import math
import re
import subprocess
import tempfile
from pathlib import Path

from .pdf import extract_pdf_region


REPORT_GUIDANCE = """
科研图表与组会汇报（适配 Mimir research-figure-plan / research-meeting-deck）：
先规划主张、真实数据来源及版本、坐标/单位和形式，再交文件。无数据仅交规划；如用户要求示意，可用 HTML 并在标题和图注明确“示意图，非实验结果”。不合成实验值。
PNG 图表调用 write_file(kind="png")，content 是 JSON：{paper_id,version_id,x,ys,style,claim}。输入必须是项目内已保存的纯CSV粘贴材料，工具直接读取真实行，禁止传values。x为横轴列名，ys为数据列名数组，style为bar/line/scatter，claim为图表目的（不是已验证的结论）。列标题保留原单位；不同单位分图。不额外百分比换算、插值或排序。缺数据请说明并请用户通过“添加来源”粘贴CSV（带列标题及单位）。
PPTX 调用 write_file(kind="pptx")，content 是 JSON：{slides:[{title,text,figure?,chart?,sources?}]}。每页文本重要结论用[cite:E1]，对应本轮citation_ids。标题简短，图页正文建议120个汉字以内；放不下会拒绝，请精简或明确拆页，不自动缩字或增加页数。figure={paper_id,version_id,page,rect?}，rect为显示页坐标，可省略以展示整页；不猜裁剪框或声称整页是已识别的Figure。工具只提取原页，不理解图意，图注须有文本或用户依据。chart={artifact_id,version_id}使用实际保存的PNG数据图，在PPT中是可编辑图表。每页至多一个figure或chart。sources=[{artifact_id,version_id}]固定所使用的项目成果版本。
已有材料即可独立做PPT，不等待新图表或日报。进展和笔记可来自粘贴材料。跨对话选材用list_files(project_wide=true)、read_file(project_wide=true)；读到旧成果中的引用不代表本轮已核对，重要事实应回查材料。sources只记录真正使用过的成果；不因打开/选择就生成文件。PPTX可只用文字；缺图片/缺结果写明，保留已有内容。warnings为真实资源缺口，交付时说明并允许补齐后追加新版；不能把资源缺失说成成功提图。
文件保存仍用write_file，修订先read_file读基线。图表数据/参数/来源版本可下载；PPT每页备注保留来源及版本。只使用返回的文件链接，不自行拼出未生成文件。
"""


def report_file(tools, ref):
    object_fields(ref, ('artifact_id', 'version_id'))
    item = tools.store.artifact(tools.project, text(ref['artifact_id'], 100))
    version = next((v for v in item['versions'] if v['id'] == ref['version_id']), None) if item else None
    if not version:
        raise ValueError('汇报成果版本不属于当前项目')
    tools.list_materials()
    if any(m['paper_id'] not in tools.allowed for m in version['materials']):
        raise ValueError('汇报成果依据超出本轮材料范围')
    return version


def render_report(tools, kind, title, content, citations):
    spec = json.loads(content)
    materials, warnings = {}, []
    def material(ref):
        paper = tools.material_version(text(ref['paper_id'], 100), text(ref['version_id'], 100))
        materials[paper['version_id']] = {'paper_id': paper['id'], 'paper_version_id': paper['version_id'], 'title': paper['title']}
        return paper
    if kind == 'png':
        object_fields(spec, ('paper_id', 'version_id', 'x', 'ys', 'style', 'claim'))
        paper = material(spec)
        if not paper['path'].startswith('paste:'):
            raise ValueError('作图需要已保存的 CSV 文字材料，不能从 PDF 猜造数据')
        csv_text = ''.join(p['text'] for p in json.loads(paper['pages']))
        if hashlib.sha256(csv_text.encode()).hexdigest() != paper['sha256']:
            raise ValueError('数据内容与材料版本不一致')
        data = chart_data(csv_text, spec)
        raw, renderer = render_chart(title, data)
        metadata = {'parameters': spec, 'csv': csv_text, 'source_sha256': paper['sha256'], 'data': data,
                    'renderer': renderer, 'warnings': [], 'summary': spec['claim']}
        return raw, content, metadata, list(materials.values())
    object_fields(spec, ('slides',))
    if not isinstance(spec['slides'], list) or not 1 <= len(spec['slides']) <= 40:
        raise ValueError('PPTX 需要 1–40 页内容')
    available = {c['id']: c for c in citations}
    slides, summaries = [], []
    for index, source in enumerate(spec['slides'], 1):
        object_fields(source, ('title', 'text'), ('figure', 'chart', 'sources'))
        heading, body = text(source['title'], 60), text(source['text'])
        used = re.findall(r'\[cite:([^\]]+)\]', heading + '\n' + body)
        if any(cid not in available for cid in used) or '[cite:' in re.sub(r'\[cite:[^\]]+\]', '', heading + body):
            raise ValueError('PPTX 引用未列入 citation_ids 或标记格式错误')
        notes = {'citations': [available[cid] for cid in dict.fromkeys(used)], 'sources': []}
        numbered = {cid: index + 1 for index, cid in enumerate(dict.fromkeys(used))}
        notes['citation_numbers'] = numbered
        item = {'title': re.sub(r'\[cite:[^\]]+\]', '', heading),
                'text': re.sub(r'\[cite:([^\]]+)\]', lambda m: f'[{numbered[m[1]]}]', body), 'provenance': notes}
        refs = source.get('sources', [])
        if not isinstance(refs, list) or len(refs) > 20:
            raise ValueError('每页最多选择20个成果来源')
        for ref in refs:
            version = report_file(tools, ref)
            if version['id'] not in tools.read_versions:
                raise ValueError('汇报选材前须 read_file 读取指定成果版本')
            notes['sources'].append({**ref, 'title': version['title'], 'materials': version['materials'], 'citations': version['citations']})
            materials.update({m['paper_version_id']: m for m in version['materials']})
        if 'figure' in source and 'chart' in source:
            raise ValueError('每页只放一张论文图片或数据图；请拆页')
        if 'figure' in source:
            ref = object_fields(source['figure'], ('paper_id', 'version_id', 'page'), ('rect',))
            paper = material(ref)  # Authorization failures must not become optional-resource warnings.
            if paper['path'].startswith('paste:'):
                raise ValueError('论文图片须来自 PDF 材料')
            try:
                raw = Path(paper['source_path']).read_bytes() if paper['source_path'] else b''
                item['image'], info = extract_pdf_region(raw, paper['sha256'], ref['page'], ref.get('rect'))
                notes['figure'] = {**ref, **info, 'title': paper['title']}
                item['text'] += f'\n原文第 {ref["page"]} 页' + ('（指定区域）' if 'rect' in ref else '（整页，非自动识别图）')
            except (OSError, ValueError, RuntimeError) as error:
                warning = f'第 {index} 页图片未提取：' + failure_message(tools.store, error, task_id=tools.task['id'], operation='report_figure')
                warnings.append(warning)
                notes['missing_figure'] = {**ref, 'reason': user_message(error)}
                item['text'] += '\n' + warning
        if 'chart' in source:
            version = report_file(tools, source['chart'])
            if version['kind'] != 'png' or 'report' not in version['payload']:
                raise ValueError('PPT 图表须使用已保存的科研 PNG 图表版本')
            materials.update({m['paper_version_id']: m for m in version['materials']})
            report = version['payload']['report']
            workspace = (tools.store.root / 'workspaces' / tools.project).resolve()
            filename = version['payload'].get('filename', '')
            path = (workspace / filename).resolve()
            if path.parent != workspace or not filename:
                raise ValueError('图表文件位置无效')
            try:
                if hashlib.sha256(path.read_bytes()).hexdigest() != version['payload']['sha256']:
                    raise ValueError('图表文件校验失败')
                item['chart'] = chart_data(report['csv'], report['parameters'])
                notes['chart'] = {**source['chart'], 'title': version['title'], **report}
            except (OSError, ValueError) as error:
                warning = f'第 {index} 页图表未载入：' + failure_message(tools.store, error, task_id=tools.task['id'], operation='report_chart')
                warnings.append(warning)
                notes['missing_chart'] = {**source['chart'], 'reason': user_message(error)}
                item['text'] += '\n' + warning
        slides.append(item)
        summaries.append({'title': item['title'], 'text': item['text'], 'provenance': notes})
    raw, count = render_deck(title, slides)
    return raw, content, {'slides': summaries, 'slide_count': count, 'warnings': warnings,
                          'renderer': 'mimir-a568b373/renderDeck+pptxgenjs-4.0.1; python-pptx-1.0.2 source/chart adapter', 'summary': f'{count} 页组会汇报'}, list(materials.values())


def object_fields(value, required, optional=()):
    if not isinstance(value, dict) or not set(required) <= value.keys() or value.keys() - set(required) - set(optional):
        raise ValueError('文件参数缺失或包含未知字段')
    return value


def text(value, limit=4000):
    if not isinstance(value, str) or not value.strip() or len(value) > limit:
        raise ValueError(f'文字须为 1–{limit} 字符')
    return value


def chart_data(csv_text, spec):
    object_fields(spec, ('paper_id', 'version_id', 'x', 'ys', 'style', 'claim'))
    text(spec['claim'], 300)
    if spec['style'] not in ('bar', 'line', 'scatter'):
        raise ValueError('图表仅支持 bar、line、scatter')
    ys = spec['ys']
    if not isinstance(ys, list) or not 1 <= len(ys) <= 8 or any(not isinstance(y, str) for y in ys) or len(set(ys)) != len(ys):
        raise ValueError('请选择 1–8 个不同的数据列')
    reader = csv.DictReader(io.StringIO(csv_text.lstrip('\ufeff')), strict=True)
    headers = reader.fieldnames
    if not headers or any(not h.strip() for h in headers) or len(set(headers)) != len(headers):
        raise ValueError('CSV 须包含非空且唯一的列标题；单位写在列标题中')
    if spec['x'] not in headers or any(y not in headers or y == spec['x'] for y in ys):
        raise ValueError('横轴和数据列必须来自 CSV 的不同列')
    rows = list(reader)
    if not rows or len(rows) > (60 if spec['style'] == 'bar' else 2000):
        raise ValueError('没有真实数据或行数过多（柱状图最多60行，其他图2000行），请明确选择数据子集')
    if any(None in row or any(v is None for v in row.values()) for row in rows):
        raise ValueError('CSV 行列数不一致，未生成图表')
    def number(value):
        try:
            result = float(value)
        except (ValueError, TypeError):
            raise ValueError('数据列必须为有限数字；缺失值不补造') from None
        if not math.isfinite(result):
            raise ValueError('数据含 NaN 或无穷值，未生成图表')
        return result
    x = [row[spec['x']] for row in rows]
    if any(not v.strip() or len(v) > 80 for v in x):
        raise ValueError('横轴存在空值或过长标签')
    if spec['style'] != 'bar':
        x = [number(v) for v in x]
        if spec['style'] == 'line' and any(a >= b for a, b in zip(x, x[1:])):
            raise ValueError('折线图横轴须严格递增；不会擅自排序或聚合数据')
    def unit(label):
        match = re.search(r'(?:\(([^()]+)\)|（([^（）]+)）|\[([^\[\]]+)\]|【([^【】]+)】)\s*$', label)
        return next(group.strip() for group in match.groups() if group is not None) if match else ''
    units = {unit(y) for y in ys}
    if len(ys) > 1 and '' in units:
        raise ValueError('多个系列须在列标题注明相同单位（无量纲也须注明），不能推测单位相同')
    if len(units) > 1:
        raise ValueError('不同单位请分图展示；支持 名称(单位)、名称[单位] 和中文括号')
    return {'x': x, 'series': [{'name': y, 'values': [number(row[y]) for row in rows]} for y in ys],
            'x_label': spec['x'], 'y_label': ys[0] if len(ys) == 1 else next(iter(units)),
            'style': spec['style'], 'claim': spec['claim']}


def call_mimir(request, directory, suffix):
    directory = Path(directory)
    source, output = directory / 'input.json', directory / f'output.{suffix}'
    source.write_text(json.dumps(request, ensure_ascii=False), encoding='utf-8')
    try:
        result = subprocess.run(['node', str(Path(__file__).with_name('mimir.mjs')), str(source), str(output)],
                                capture_output=True, timeout=60)
    except (OSError, subprocess.TimeoutExpired) as error:
        raise ValueError('MethodAtlas 图表渲染器不可用；请安装 Node.js >=22.18 并在仓库运行 npm ci') from error
    if result.returncode == 2 and request['kind'] == 'bar':
        return None  # Known upstream formatting limit, not an execution failure.
    if result.returncode:
        raise ValueError('MethodAtlas 图表渲染失败：' + result.stderr.decode('utf-8', errors='replace')[-1500:])
    return output.read_bytes()


def render_chart(title, data):
    from PIL import ImageFont
    from matplotlib.font_manager import findfont
    caption = f'{title} · {data["x_label"]} / {data["series"][0]["name"]}'
    font = ImageFont.truetype(findfont('Microsoft YaHei'), 15)
    if data['style'] == 'bar' and len(data['series']) == 1 and font.getlength(caption) <= 600:
        rows = [{'name': str(x), 'value': y} for x, y in zip(data['x'], data['series'][0]['values'])]
        with tempfile.TemporaryDirectory() as directory:
            raw = call_mimir({'kind': 'bar', 'title': caption, 'rows': rows}, directory, 'png')
        if raw is not None:
            return raw, 'mimir-a568b373/metricFigureSvg+resvg-2.6.2'
    from matplotlib import rc_context
    from matplotlib.figure import Figure
    with rc_context({'font.family': ['Microsoft YaHei', 'DejaVu Sans'], 'axes.unicode_minus': False}):
        figure = Figure(figsize=(10, 6), layout='constrained')
        axis = figure.subplots()
        for index, series in enumerate(data['series']):
            if data['style'] == 'bar':
                width = .8 / len(data['series'])
                bars = axis.bar([i - .4 + width / 2 + index * width for i in range(len(data['x']))], series['values'], width, label=series['name'])
                if len(data['x']) * len(data['series']) <= 30:
                    axis.bar_label(bars, fmt='%g', padding=3)
                axis.set_xticks(range(len(data['x'])), data['x'], rotation=30 if len(data['x']) > 6 else 0)
            else:
                axis.plot(data['x'], series['values'], marker=['o', 's', '^', 'D', 'v', 'P', 'X', '*'][index],
                          linestyle='-' if data['style'] == 'line' else 'None', label=series['name'])
        axis.set(title=title, xlabel=data['x_label'], ylabel=data['y_label'])
        axis.legend()
        axis.grid(axis='y', alpha=.2)
        output = io.BytesIO()
        figure.savefig(output, format='png', dpi=160, metadata={'Description': json.dumps(data, ensure_ascii=False)})
        return output.getvalue(), 'matplotlib-3.10.8 (Mimir format not supported)'


def wrap_lines(value, width, font):
    """Measure the installed font; avoid Office wrapping pre-wrapped lines again."""
    lines = []
    for paragraph in value.splitlines() or ['']:
        line = ''
        for char in re.findall(r'\d+(?:\.\d+)?|[A-Za-z][A-Za-z0-9_-]*|.', paragraph):
            if font.getlength(char) > width:
                raise ValueError('单词或标识过长，请将完整标识放入来源备注，正文改用简短名称')
            if line and font.getlength(line + char) > width:
                lines.append(line)
                line = ''
            line += char
        lines.append(line)
    return lines


def render_deck(title, slides):
    """Mimir owns slide layout; adapt validated inputs, notes and native charts."""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData, XyChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Inches, Pt
    from PIL import Image, ImageFont
    from matplotlib.font_manager import findfont
    font_path = findfont('Microsoft YaHei')
    footer_font = ImageFont.truetype(font_path, 8 * 4)
    footer = title.replace('\n', ' ')
    if footer_font.getlength(footer) > 6 * 72 * 4:
        while footer and footer_font.getlength(footer + '…') > 6 * 72 * 4:
            footer = footer[:-1]
        footer += '…'
    plans = []
    with tempfile.TemporaryDirectory() as directory:
        for index, item in enumerate(slides):
            visual = item.get('image') is not None or item.get('chart') is not None
            size, width, max_lines = (12.5, 2.4, 12) if visual else (14, 8, 10)
            lines = wrap_lines(item['text'], width * 72 * 4, ImageFont.truetype(font_path, round(size * 4)))
            headings = wrap_lines(item['title'], 8.6 * 72 * 4, ImageFont.truetype(font_path, 21 * 4))
            if len(lines) > max_lines or len(headings) > 1:
                raise ValueError(f'幻灯片“{item["title"]}”文字过长；请精简正文或明确拆页')
            if visual:
                path = Path(directory) / f'{index}.png'
                if item.get('image') is not None:
                    path.write_bytes(item['image'])
                else:
                    # Only a local placeholder: replaced by an editable chart below.
                    Image.new('RGB', (10, 10), 'white').save(path)
                plans.append({'kind': 'figure', 'heading': item['title'], 'imagePath': str(path), 'caption': '\n'.join(lines)})
            else:
                plans.append({'kind': 'bullets', 'heading': item['title'], 'bullets': [{'text': line} for line in lines]})
        raw = call_mimir({'kind': 'pptx', 'title': footer, 'slides': plans}, directory, 'pptx')
    presentation = Presentation(io.BytesIO(raw))
    presentation.core_properties.title = title
    for slide, item in zip(presentation.slides, slides):
        # Upstream has no provenance/native-chart fields. Its image sizing option
        # also needs aspect-ratio correction with pptxgenjs 4.0.1.
        for shape in list(slide.shapes):
            if shape.has_text_frame:
                frame = shape.text_frame
                frame.word_wrap = False
                frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
            if shape.shape_type == 13:
                if item.get('chart') is not None:
                    shape._element.getparent().remove(shape._element)
                else:
                    with Image.open(io.BytesIO(item['image'])) as image:
                        w, h = image.size
                    scale = min(5.54 / w, 3.44 / h)
                    shape.left, shape.top = Inches(.78 + (5.54 - w * scale) / 2), Inches(1.53 + (3.44 - h * scale) / 2)
                    shape.width, shape.height = Inches(w * scale), Inches(h * scale)
        if item.get('chart') is not None:
            data = item['chart']
            values = CategoryChartData() if data['style'] == 'bar' else XyChartData()
            if data['style'] == 'bar':
                values.categories = data['x']
            for series in data['series']:
                if data['style'] == 'bar':
                    values.add_series(series['name'], series['values'])
                else:
                    points = values.add_series(series['name'])
                    for x, y in zip(data['x'], series['values']):
                        points.add_data_point(x, y)
            kind = {'bar': XL_CHART_TYPE.COLUMN_CLUSTERED, 'line': XL_CHART_TYPE.XY_SCATTER_LINES,
                    'scatter': XL_CHART_TYPE.XY_SCATTER}[data['style']]
            chart = slide.shapes.add_chart(kind, Inches(.78), Inches(1.53), Inches(5.54), Inches(3.44), values).chart
            chart.has_title = False
            chart.plots[0].vary_by_categories = False
            chart.has_legend = len(data['series']) > 1
            if chart.has_legend:
                chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                chart.legend.font.size = Pt(10)
            for axis, label in [(chart.category_axis, data['x_label']), (chart.value_axis, data['y_label'])]:
                axis.has_title = True
                axis.axis_title.text_frame.text = label
                axis.axis_title.text_frame.paragraphs[0].font.size = Pt(11)
                axis.tick_labels.font.size = Pt(10)
            if data['style'] == 'bar':
                chart.value_axis.minimum_scale = min(0, min(v for s in data['series'] for v in s['values']))
        slide.notes_slide.notes_text_frame.text = json.dumps(item['provenance'], ensure_ascii=False, indent=2)
    output = io.BytesIO()
    presentation.save(output)
    raw = output.getvalue()
    Presentation(io.BytesIO(raw))
    return raw, len(presentation.slides)
